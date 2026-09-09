import os
import re
import math
import datetime
import frappe
from frappe import _
from frappe.utils.csvutils import read_csv_content
from frappe.utils import flt, get_datetime_str, now_datetime, getdate, today


FIELD_MAP = {
    # Employee Number / Code / ID
    "employee_number": "employee_number",
    "emp_number": "employee_number",
    "emp_code": "employee_number",
    "emp_id": "employee_number",
    "employee_code": "employee_number",
    "employee_id": "employee_number",
    "emp_no": "employee_number",
    "employee_no": "employee_number",
    "staff_id": "employee_number",
    "user_id": "user_id",

    # Names
    "first_name": "first_name",
    "middle_name": "middle_name",
    "last_name": "last_name",
    "employee_name": "employee_name",
    "emp_name": "employee_name",
    "name": "employee_name",
    "full_name": "employee_name",

    # Demographics & Dates
    "gender": "gender",
    "date_of_birth": "date_of_birth",
    "dob": "date_of_birth",
    "birth_date": "date_of_birth",
    "date_of_joining": "date_of_joining",
    "doj": "date_of_joining",
    "joining_date": "date_of_joining",
    "final_confirmation_date": "final_confirmation_date",
    "confirmation_date": "final_confirmation_date",
    "doc": "final_confirmation_date",
    "relieving_date": "relieving_date",
    "date_of_relieving": "relieving_date",
    "dor": "relieving_date",
    "resignation_letter_date": "resignation_letter_date",
    "resignation_date": "resignation_letter_date",

    # Organization / Designations
    "designation": "designation",
    "desig": "designation",
    "department": "department",
    "dept": "department",
    "sub_department": "sub_department",
    "subdepartment": "sub_department",
    "sub_dept": "sub_department",
    "subdept": "sub_department",
    "branch": "branch",
    "branch_name": "branch",
    "sol_id": "sol_id",
    "sol": "sol_id",
    "sahayog_branch": "sahayog_branch",
    "sub_branch": "sahayog_branch",
    "sahayog_branch_name": "sahayog_branch",

    # Contact & Personal
    "mobile_number": "cell_number",
    "mobile_no": "cell_number",
    "mobileno": "cell_number",
    "mobile": "cell_number",
    "phone": "cell_number",
    "phone_number": "cell_number",
    "cell": "cell_number",
    "cell_number": "cell_number",
    "contact_no": "cell_number",
    "contact": "cell_number",
    "personal_email": "personal_email",
    "email": "personal_email",
    "email_id": "personal_email",
    "company": "company",

    # Financial / Identity
    "bank_name": "bank_name",
    "bank_account_number": "bank_ac_no",
    "bank_ac_no": "bank_ac_no",
    "bank_account": "bank_ac_no",
    "account_no": "bank_ac_no",
    "account_number": "bank_ac_no",
    "marital_status": "marital_status",
    "marital": "marital_status",
    "blood_group": "blood_group",
    "blood": "blood_group",
    "permanent_address": "permanent_address",
    "current_address": "current_address",
    "shift": "default_shift",
    "default_shift": "default_shift",
    "employment_type": "employment_type",
    "typeofemployment": "employment_type",
    "reports_to": "reports_to",
    "reporting_to": "reports_to",
    "reportingmanager_code": "reports_to",
    "manager": "reports_to",
    "manager_code": "reports_to",
    "status": "status",
    "employee_status": "status",
    "pan_number": "pan_number",
    "pan": "pan_number",
    "pan_no": "pan_number",
    "aadhaar_card_number": "aadhaar_card_number",
    "custom_aadhar_number": "custom_aadhar_number",
    "adharcard_no": "custom_aadhar_number",
    "aadhar": "aadhaar_card_number",
    "aadhaar": "aadhaar_card_number",
    "aadhaar_no": "aadhaar_card_number",
    "aadhar_no": "aadhaar_card_number",
    "uhid_number": "custom_uhid_number",
    "custom_uhid_number": "custom_uhid_number",
    "uhid": "custom_uhid_number",
    "monthly_gross_salary": "ctc",
    "monthly_salary": "ctc",
    "gross_salary": "ctc",
    "salary": "ctc",
    "ctc": "ctc",

    # Custom Org Hierarchy
    "zone": "custom_zone",
    "custom_zone": "custom_zone",
    "region": "custom_region",
    "custom_region": "custom_region",
    "statename": "custom_state",
    "district": "custom_district",
    "districtname": "custom_district",
    "custom_district": "custom_district",
    "division": "custom_division",
    "custom_division": "custom_division",
}

DEFAULT_MANDATORY = [
    "first_name",
    "gender",
    "date_of_joining",
    "designation",
    "department",
]

EMP_NUM_ALIASES = (
    "employee_number", "emp_number", "emp_code", "emp_id",
    "employee_code", "emp_no", "employee_no", "staff_id", "user_id"
)

FIELD_ALIASES = {
    "employee_number": EMP_NUM_ALIASES,
    "first_name": ("first_name", "employee_name", "emp_name", "name", "full_name"),
    "date_of_joining": ("date_of_joining", "doj", "joining_date"),
    "designation": ("designation", "desig"),
    "department": ("department", "dept"),
    "sub_department": ("sub_department", "subdepartment", "sub_dept", "subdept", "sub department"),
    "gender": ("gender", "sex"),
    "cell_number": ("cell_number", "mobile_number", "mobile_no", "mobileno", "mobile", "phone", "phone_number", "contact_no"),
    "date_of_birth": ("date_of_birth", "dob", "birth_date"),
    "relieving_date": ("relieving_date", "date_of_relieving", "dor"),
}

CACHE_KEY = "emp_import_session:{user}"


def _clean_header(h):
    if not h:
        return ""
    h = str(h).lstrip("\ufeff").strip().lower()
    h = re.sub(r"[\s\-_.]+", "_", h)
    return h.strip("_")


def _clean_val(val):
    if val is None:
        return None
    if isinstance(val, (int,)):
        return str(val)
    if isinstance(val, float):
        if math.isnan(val):
            return None
        if val.is_integer():
            return str(int(val))
        return f"{val:.10f}".rstrip("0").rstrip(".")
    val_str = str(val).strip()
    if val_str.endswith(".0") and val_str[:-2].isdigit():
        return val_str[:-2]
    return val_str


def _extract_emp_number(row_dict, header_for):
    emp_key = header_for.get("employee_number", "employee_number")
    raw = row_dict.get(emp_key)
    if raw is not None and str(raw).strip():
        return str(raw).strip()

    for alias in EMP_NUM_ALIASES:
        v = row_dict.get(alias)
        if v is not None and str(v).strip():
            return str(v).strip()
    return ""


def _has_field_value(row_dict, field_name, header_for):
    csv_key = header_for.get(field_name, field_name)
    val = row_dict.get(csv_key)
    if val is not None and str(val).strip():
        return True

    aliases = FIELD_ALIASES.get(field_name, ())
    for a in aliases:
        v = row_dict.get(a)
        if v is not None and str(v).strip():
            return True
    return False


def _get_or_build_session_cache(setting, force=False):
    key = CACHE_KEY.format(user=frappe.session.user)
    cached = frappe.cache().get_value(key)
    file_url = setting.get("employee_master")

    if cached and not force:
        if cached.get("file_url") == file_url:
            return cached

    rows = _parse_file(file_url)
    if not rows or len(rows) < 2:
        frappe.throw(_("File has no valid data rows"))

    lookup_cache = _load_lookup_cache()
    existing_cols = list(r[0] for r in frappe.db.sql("SHOW COLUMNS FROM `tabEmployee`"))

    data = {
        "rows": rows,
        "lookup_cache": lookup_cache,
        "existing_cols": existing_cols,
        "file_url": file_url,
    }
    frappe.cache().set_value(key, data, expires_in_sec=3600)
    return data


@frappe.whitelist()
def init_import_session(mode="insert", batch_size=100):
    setting = frappe.get_doc("Sahayog HR Setting")
    file_url = setting.get("employee_master")
    if not file_url:
        frappe.throw(_("Please upload an Employee Master file first"))

    cached_data = _get_or_build_session_cache(setting, force=True)
    rows = cached_data["rows"]

    total_rows = len(rows) - 1
    batch_size = max(1, int(batch_size))
    total_batches = math.ceil(total_rows / batch_size)

    return {
        "total_rows": total_rows,
        "batch_size": batch_size,
        "total_batches": total_batches,
    }


@frappe.whitelist()
def process_import_batch(mode="insert", batch_index=0, batch_size=100):
    setting = frappe.get_doc("Sahayog HR Setting")
    file_url = setting.get("employee_master")
    if not file_url:
        frappe.throw(_("Please upload an Employee Master file first"))

    cached_data = _get_or_build_session_cache(setting)
    rows = cached_data["rows"]
    lookup_cache = cached_data["lookup_cache"]
    existing_cols = cached_data["existing_cols"]

    table_mappings = _load_table_mappings(setting)
    mandatory_fields = _get_mandatory_fields(table_mappings)

    raw_headers = rows[0]
    headers = [_clean_header(h) for h in raw_headers]
    data_rows = rows[1:]

    valid_field_map = _build_field_map(headers, existing_cols, table_mappings)
    header_for = {v: k for k, v in valid_field_map.items()}

    batch_index = int(batch_index)
    batch_size = int(batch_size)

    start_idx = batch_index * batch_size
    end_idx = min((batch_index + 1) * batch_size, len(data_rows))
    batch_rows = data_rows[start_idx:end_idx]

    result = {
        "inserted": 0,
        "updated": 0,
        "skipped": 0,
        "failed": 0,
        "errors": [],
        "inserted_numbers": [],
        "updated_numbers": [],
    }

    frappe.flags.in_import = True

    for offset, row in enumerate(batch_rows):
        i = start_idx + offset + 2
        row_dict = _row_to_dict(headers, row)

        emp_number = _extract_emp_number(row_dict, header_for)

        # Ignore footer / watermark / total rows
        if emp_number and emp_number.lower() in ("super employee master", "total", "grand total", "summary", "end of report"):
            result["skipped"] += 1
            continue

        # Ignore rows with fewer than 2 non-empty fields (likely metadata or blank lines)
        non_empty_count = sum(1 for v in row if v is not None and str(v).strip() != "")
        if non_empty_count < 2:
            result["skipped"] += 1
            continue

        first_name_key = header_for.get("first_name", "first_name")
        raw_fn = row_dict.get(first_name_key) or row_dict.get("employee_name") or row_dict.get("emp_name") or row_dict.get("name")
        emp_label = str(raw_fn).strip() if raw_fn is not None else f"Row {i}"

        if not emp_number:
            result["failed"] += 1
            result["errors"].append(f"Row {i}: {emp_label} - employee_number is missing")
            continue

        # Mandatory field check with smart alias support
        if mode == "insert" or table_mappings:
            missing_mandatory = []
            for f in mandatory_fields:
                if not _has_field_value(row_dict, f, header_for):
                    missing_mandatory.append(f)
            if missing_mandatory:
                result["failed"] += 1
                result["errors"].append(f"Row {i}: {emp_number} - Required field(s) missing: {', '.join(missing_mandatory)}")
                continue

        existing_emp_name = lookup_cache["Employee"].get(emp_number) or lookup_cache["Employee"].get(emp_number.lower())
        if not existing_emp_name:
            existing_emp_name = (
                frappe.db.exists("Employee", {"employee_number": emp_number})
                or frappe.db.exists("Employee", emp_number)
            )
            if existing_emp_name:
                lookup_cache["Employee"][emp_number] = existing_emp_name

        # Process single row within a savepoint for complete error isolation
        sp_name = f"sp_imp_{offset}"
        try:
            frappe.db.savepoint(sp_name)

            if mode == "insert":
                if existing_emp_name:
                    result["skipped"] += 1
                else:
                    emp_name = _create_employee(
                        row_dict,
                        valid_field_map,
                        emp_number=emp_number,
                        cache=lookup_cache,
                        existing_cols=existing_cols,
                    )
                    result["inserted"] += 1
                    result["inserted_numbers"].append(emp_number)
                    lookup_cache["Employee"][emp_number] = emp_name
                    lookup_cache["Employee"][emp_name] = emp_name

            elif mode == "update":
                if not existing_emp_name:
                    result["skipped"] += 1
                else:
                    _update_employee(
                        existing_emp_name,
                        row_dict,
                        valid_field_map,
                        cache=lookup_cache,
                        existing_cols=existing_cols,
                    )
                    result["updated"] += 1
                    result["updated_numbers"].append(emp_number)

            frappe.db.release_savepoint(sp_name)

        except _StopRow:
            frappe.db.rollback(save_point=sp_name)
            continue
        except frappe.DuplicateEntryError:
            frappe.db.rollback(save_point=sp_name)
            if mode == "insert":
                result["skipped"] += 1
                lookup_cache["Employee"][emp_number] = emp_number
            else:
                result["failed"] += 1
                result["errors"].append(f"Row {i}: {emp_number} - Duplicate entry error")
        except Exception as e:
            frappe.db.rollback(save_point=sp_name)
            err_str = str(e).strip()
            if "Duplicate entry" in err_str:
                result["skipped"] += 1
                lookup_cache["Employee"][emp_number] = emp_number
            else:
                result["failed"] += 1
                result["errors"].append(f"Row {i}: {emp_number or emp_label} - {err_str}")

    frappe.db.commit()
    frappe.flags.in_import = False

    key = CACHE_KEY.format(user=frappe.session.user)
    cached_data["lookup_cache"] = lookup_cache
    frappe.cache().set_value(key, cached_data, expires_in_sec=3600)

    return result


@frappe.whitelist()
def finish_import_session(summary_data, mode="insert"):
    if isinstance(summary_data, str):
        import json
        summary_data = json.loads(summary_data)

    summary = _build_summary(summary_data, mode)
    frappe.db.set_value("Sahayog HR Setting", None, "employee_import_summary", summary)
    frappe.db.commit()

    return summary


@frappe.whitelist()
def import_employee_master(mode="insert"):
    init_res = init_import_session(mode=mode, batch_size=100)
    total_batches = init_res["total_batches"]

    aggregated = {
        "inserted": 0,
        "updated": 0,
        "skipped": 0,
        "failed": 0,
        "errors": [],
        "inserted_numbers": [],
        "updated_numbers": [],
    }

    for b in range(total_batches):
        batch_res = process_import_batch(mode=mode, batch_index=b, batch_size=100)
        aggregated["inserted"] += batch_res.get("inserted", 0)
        aggregated["updated"] += batch_res.get("updated", 0)
        aggregated["skipped"] += batch_res.get("skipped", 0)
        aggregated["failed"] += batch_res.get("failed", 0)
        aggregated["errors"].extend(batch_res.get("errors", []))
        aggregated["inserted_numbers"].extend(batch_res.get("inserted_numbers", []))
        aggregated["updated_numbers"].extend(batch_res.get("updated_numbers", []))

    return finish_import_session(aggregated, mode=mode)


def _resolve_filepath(file_url):
    site_path = os.path.abspath(frappe.get_site_path())
    if file_url.startswith("/private/"):
        return site_path + file_url
    return site_path + "/public" + file_url


def _parse_file(file_url):
    clean_url = file_url.split("?")[0].lower()

    if clean_url.endswith(".xlsx") or clean_url.endswith(".xls") or clean_url.endswith(".xlsm"):
        return _parse_excel(_resolve_filepath(file_url))

    elif clean_url.endswith(".csv"):
        return _parse_csv(file_url)

    else:
        try:
            return _parse_excel(_resolve_filepath(file_url))
        except Exception:
            return _parse_csv(file_url)


def _parse_excel(filepath):
    import openpyxl

    wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
    rows = []

    sheets_to_check = [wb.active] if wb.active else []
    for sname in wb.sheetnames:
        s = wb[sname]
        if s not in sheets_to_check:
            sheets_to_check.append(s)

    for sheet in sheets_to_check:
        sheet_rows = []
        for r in sheet.iter_rows(values_only=True):
            if any(cell is not None and str(cell).strip() != "" for cell in r):
                sheet_rows.append(list(r))
        if len(sheet_rows) > len(rows):
            rows = sheet_rows

    wb.close()
    return rows


def _parse_csv(file_url):
    filepath = _resolve_filepath(file_url)
    content = None

    try:
        file_doc = frappe.get_doc("File", {"file_url": file_url})
        content = file_doc.get_content()
    except Exception:
        pass

    if content is None:
        with open(filepath, "rb") as f:
            content = f.read()

    if isinstance(content, bytes):
        for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            text = content.decode("utf-8", errors="ignore")
    else:
        text = str(content)

    raw_rows = read_csv_content(text) or []
    clean_rows = []
    for r in raw_rows:
        if any(cell is not None and str(cell).strip() != "" for cell in r):
            clean_rows.append(r)

    return clean_rows


def _load_table_mappings(setting):
    if not setting.get("field_mappings"):
        return None
    mappings = []
    for row in setting.field_mappings:
        if not row.enabled:
            continue
        mappings.append({
            "source_column": row.source_column,
            "target_field": row.target_field,
            "is_mandatory": row.is_mandatory,
        })
    return mappings


def _get_mandatory_fields(table_mappings):
    if table_mappings:
        return [m["source_column"] for m in table_mappings if m.get("is_mandatory")]
    return DEFAULT_MANDATORY


def _build_field_map(headers, existing_cols, table_mappings=None):
    mapping = {}
    clean_headers = [_clean_header(h) for h in headers]

    if table_mappings:
        for m in table_mappings:
            src_clean = _clean_header(m.get("source_column", ""))
            tgt = m.get("target_field", "").strip()
            if src_clean and tgt and src_clean in clean_headers:
                if tgt in existing_cols:
                    mapping[src_clean] = tgt
                elif f"custom_{tgt}" in existing_cols:
                    mapping[src_clean] = f"custom_{tgt}"

    for h_raw in headers:
        clean = _clean_header(h_raw)
        if clean and clean not in mapping:
            doc_field = FIELD_MAP.get(clean, clean)
            if doc_field in existing_cols:
                mapping[clean] = doc_field
            elif f"custom_{doc_field}" in existing_cols:
                mapping[clean] = f"custom_{doc_field}"

    return mapping


def _row_to_dict(headers, row):
    d = {}
    for idx, h in enumerate(headers):
        clean_h = _clean_header(h)
        val = row[idx] if idx < len(row) else None
        cleaned_val = _clean_val(val)
        d[clean_h] = cleaned_val
    return d


def _load_lookup_cache():
    cache = {
        "Division": {},
        "Zone": {},
        "Region": {},
        "Branch": {},
        "Designation": {},
        "Department": {},
        "Sahayog Branch": {},
        "Employee": {},
    }

    for d in frappe.db.get_all("Division", fields=["name", "division"]):
        if d.division:
            cache["Division"][d.division.strip().title()] = d.name
            cache["Division"][d.division.strip().lower()] = d.name
        cache["Division"][d.name.strip().lower()] = d.name
        cache["Division"][d.name] = d.name

    for z in frappe.db.get_all("Zone", fields=["name", "zone"]):
        if z.zone:
            cache["Zone"][z.zone.strip().upper()] = z.name
            cache["Zone"][z.zone.strip().lower()] = z.name
        cache["Zone"][z.name.strip().upper()] = z.name
        cache["Zone"][z.name.strip().lower()] = z.name
        cache["Zone"][z.name] = z.name

    for r in frappe.db.get_all("Region", fields=["name", "region"]):
        if r.region:
            cache["Region"][r.region.strip().upper()] = r.name
            cache["Region"][r.region.strip().lower()] = r.name
        cache["Region"][r.name.strip().upper()] = r.name
        cache["Region"][r.name.strip().lower()] = r.name
        cache["Region"][r.name] = r.name

    for b in frappe.db.get_all("Branch", fields=["name", "branch"]):
        if b.branch:
            cache["Branch"][b.branch.strip().title()] = b.name
            cache["Branch"][b.branch.strip().lower()] = b.name
        cache["Branch"][b.name.strip().lower()] = b.name
        cache["Branch"][b.name] = b.name

    for d in frappe.db.get_all("Designation", fields=["name", "designation_name"]):
        if d.designation_name:
            cache["Designation"][d.designation_name.strip().title()] = d.name
            cache["Designation"][d.designation_name.strip().lower()] = d.name
        cache["Designation"][d.name.strip().lower()] = d.name
        cache["Designation"][d.name] = d.name

    for dep in frappe.db.get_all("Department", fields=["name", "department_name"]):
        if dep.department_name:
            cache["Department"][dep.department_name.strip().title()] = dep.name
            cache["Department"][dep.department_name.strip().lower()] = dep.name
        cache["Department"][dep.name.strip().lower()] = dep.name
        cache["Department"][dep.name] = dep.name

    for sb in frappe.db.get_all("Sahayog Branch", fields=["name", "zone", "region", "district", "sol_id"]):
        clean_key = sb.name.replace(" ", "")
        cache["Sahayog Branch"][clean_key] = sb
        cache["Sahayog Branch"][clean_key.lower()] = sb
        cache["Sahayog Branch"][sb.name] = sb
        if sb.get("sol_id"):
            clean_sol = str(sb["sol_id"]).strip().replace(" ", "")
            cache["Sahayog Branch"][clean_sol] = sb
            cache["Sahayog Branch"][clean_sol.lower()] = sb

    for emp in frappe.db.get_all("Employee", fields=["name", "employee_number"]):
        cache["Employee"][emp.name.strip()] = emp.name
        cache["Employee"][emp.name.strip().lower()] = emp.name
        if emp.employee_number:
            cache["Employee"][emp.employee_number.strip()] = emp.name
            cache["Employee"][emp.employee_number.strip().lower()] = emp.name

    return cache


def _ensure_link(val, target_doctype, label_field, name_prefix=None, cache=None):
    if not val:
        return val
    val = str(val).strip()
    if not val or val.lower() in ("none", "null", "na", "n/a", "-"):
        return None

    if name_prefix:
        clean = val.replace(" ", "").upper()
        if not clean.startswith(name_prefix):
            clean = name_prefix + clean
    else:
        clean = val.strip().title()

    clean_lower = clean.lower()

    if cache and target_doctype in cache:
        if clean in cache[target_doctype]:
            return cache[target_doctype][clean]
        if clean_lower in cache[target_doctype]:
            return cache[target_doctype][clean_lower]

    existing = (
        frappe.db.get_value(target_doctype, {label_field: clean}, "name")
        or frappe.db.get_value(target_doctype, clean, "name")
    )
    if not existing and not frappe.db.exists(target_doctype, clean):
        try:
            new_doc = frappe.get_doc({
                "doctype": target_doctype,
                label_field: clean,
            }).insert(ignore_permissions=True, ignore_mandatory=True)
            existing = new_doc.name
        except Exception:
            existing = clean

    res_name = existing or clean
    if cache and target_doctype in cache:
        cache[target_doctype][clean] = res_name
        cache[target_doctype][clean_lower] = res_name
        cache[target_doctype][res_name] = res_name
    return res_name


def _ensure_department(dep_val, cache=None):
    if not dep_val:
        return dep_val
    dep_val = str(dep_val).strip()
    if not dep_val or dep_val.lower() in ("none", "null", "na", "n/a", "-"):
        return None

    title_dep = dep_val.strip().title()
    lower_dep = dep_val.strip().lower()

    if cache and "Department" in cache:
        if title_dep in cache["Department"]:
            return cache["Department"][title_dep]
        if lower_dep in cache["Department"]:
            return cache["Department"][lower_dep]
        if dep_val in cache["Department"]:
            return cache["Department"][dep_val]

    existing = (
        frappe.db.get_value("Department", {"department_name": title_dep}, "name")
        or frappe.db.get_value("Department", {"department_name": dep_val}, "name")
        or frappe.db.get_value("Department", dep_val, "name")
    )
    if not existing:
        company = frappe.defaults.get_global_default("company") or frappe.db.get_value("Company", {}, "name")
        try:
            new_doc = frappe.get_doc({
                "doctype": "Department",
                "department_name": title_dep,
                "company": company,
            }).insert(ignore_permissions=True, ignore_mandatory=True)
            existing = new_doc.name
        except Exception:
            existing = frappe.db.get_value("Department", {"company": company}, "name") or title_dep

    if cache and "Department" in cache:
        cache["Department"][title_dep] = existing
        cache["Department"][lower_dep] = existing
        cache["Department"][dep_val] = existing
    return existing


def _split_name(first_name, middle_name=None, last_name=None):
    if not first_name:
        return first_name, middle_name, last_name
    parts = [p for p in str(first_name).strip().split() if p]
    if len(parts) == 1:
        return parts[0], middle_name, last_name
    if middle_name or last_name:
        return first_name, middle_name, last_name
    if len(parts) == 2:
        return parts[0], None, parts[1]
    if len(parts) == 3:
        return parts[0], parts[1], parts[2]
    return parts[0], " ".join(parts[1:-1]), parts[-1]


def _create_employee(row_dict, field_map, emp_number=None, cache=None, existing_cols=None):
    if existing_cols is None:
        existing_cols = set(r[0] for r in frappe.db.sql("SHOW COLUMNS FROM `tabEmployee`"))
    else:
        existing_cols = set(existing_cols)

    parsed = _prepare_employee_data(row_dict, field_map)

    # Ensure employee_number & name are explicitly set
    if emp_number:
        clean_num = str(emp_number).strip()
        parsed["employee_number"] = clean_num
        parsed["name"] = clean_num

    # Extract and split name safely with fallbacks
    fn = (
        parsed.get("first_name")
        or row_dict.get("first_name")
        or row_dict.get("employee_name")
        or row_dict.get("emp_name")
        or row_dict.get("name")
        or row_dict.get("full_name")
    )
    mn = parsed.get("middle_name") or row_dict.get("middle_name")
    ln = parsed.get("last_name") or row_dict.get("last_name")
    first_name, middle_name, last_name = _split_name(fn, mn, ln)
    if first_name:
        parsed["first_name"] = first_name
    if middle_name:
        parsed["middle_name"] = middle_name
    elif "middle_name" in parsed:
        del parsed["middle_name"]
    if last_name:
        parsed["last_name"] = last_name
    elif "last_name" in parsed:
        del parsed["last_name"]

    # Resolve Organization Hierarchy Links
    link_fields = [
        ("custom_division", "Division", "division", None),
        ("custom_zone", "Zone", "zone", "ZONE-"),
        ("custom_region", "Region", "region", "REGION-"),
        ("branch", "Branch", "branch", None),
        ("designation", "Designation", "designation_name", None),
    ]
    for field, doctype, label, prefix in link_fields:
        val = parsed.get(field) or row_dict.get(field)
        if val and field in existing_cols:
            parsed[field] = _ensure_link(val, doctype, label, prefix, cache=cache)

    # Resolve SOL ID and Sahayog Branch BEFORE insert
    sol_id = row_dict.get("sol_id") or row_dict.get("sol")
    branch_info = None
    if sol_id:
        clean_sol = str(sol_id).strip().replace(" ", "")
        if "sol_id" in existing_cols:
            parsed["sol_id"] = clean_sol
        if cache and "Sahayog Branch" in cache:
            branch_info = cache["Sahayog Branch"].get(clean_sol) or cache["Sahayog Branch"].get(clean_sol.lower())
        if not branch_info:
            b_list = frappe.db.sql(
                """SELECT name, zone, region, district FROM `tabSahayog Branch`
                   WHERE REPLACE(name, ' ', '') = %s OR sol_id = %s LIMIT 1""",
                (clean_sol, clean_sol),
                as_dict=True,
            )
            if b_list:
                branch_info = b_list[0]
                if cache and "Sahayog Branch" in cache:
                    cache["Sahayog Branch"][clean_sol] = branch_info

    sb_val = parsed.get("sahayog_branch") or row_dict.get("sahayog_branch") or row_dict.get("sub_branch")
    if not branch_info and sb_val:
        clean_sb = str(sb_val).strip().replace(" ", "")
        if cache and "Sahayog Branch" in cache:
            branch_info = cache["Sahayog Branch"].get(clean_sb) or cache["Sahayog Branch"].get(clean_sb.lower())
        if not branch_info:
            b_list = frappe.db.sql(
                """SELECT name, zone, region, district FROM `tabSahayog Branch`
                   WHERE REPLACE(name, ' ', '') = %s OR name = %s LIMIT 1""",
                (clean_sb, str(sb_val).strip()),
                as_dict=True,
            )
            if b_list:
                branch_info = b_list[0]
                if cache and "Sahayog Branch" in cache:
                    cache["Sahayog Branch"][clean_sb] = branch_info

    if branch_info:
        b_name = branch_info.get("name") if isinstance(branch_info, dict) else branch_info
        if "sahayog_branch" in existing_cols:
            parsed["sahayog_branch"] = b_name
        if isinstance(branch_info, dict):
            if branch_info.get("zone") and not parsed.get("custom_zone") and "custom_zone" in existing_cols:
                parsed["custom_zone"] = _ensure_link(branch_info["zone"], "Zone", "zone", "ZONE-", cache=cache)
            if branch_info.get("region") and not parsed.get("custom_region") and "custom_region" in existing_cols:
                parsed["custom_region"] = _ensure_link(branch_info["region"], "Region", "region", "REGION-", cache=cache)
            if branch_info.get("district") and not parsed.get("custom_district") and "custom_district" in existing_cols:
                parsed["custom_district"] = branch_info["district"]

    # Resolve Department
    dep_val = parsed.get("department") or row_dict.get("department") or row_dict.get("dept")
    if dep_val and "department" in existing_cols:
        parsed["department"] = _ensure_department(dep_val, cache=cache)

    # Resolve Reports To
    rt_val = parsed.get("reports_to") or row_dict.get("reports_to") or row_dict.get("reporting_to") or row_dict.get("manager")
    if rt_val and "reports_to" in existing_cols:
        rt_clean = _clean_val(rt_val)
        emp_lead = None
        if cache and "Employee" in cache:
            emp_lead = cache["Employee"].get(rt_clean) or cache["Employee"].get(str(rt_clean).strip().lower())
        if not emp_lead:
            emp_lead = (
                rt_clean if frappe.db.exists("Employee", rt_clean)
                else frappe.db.get_value("Employee", {"employee_number": rt_clean}, "name")
            )
            if emp_lead and cache and "Employee" in cache:
                cache["Employee"][rt_clean] = emp_lead
        if emp_lead:
            parsed["reports_to"] = emp_lead
        else:
            parsed.pop("reports_to", None)

    # Business Flags & Rules
    if "custom_is_support_staff" in existing_cols:
        parsed["custom_is_support_staff"] = 0
    if "custom_medical_deduction" in existing_cols and "custom_medical_deduction" not in parsed:
        parsed["custom_medical_deduction"] = 100
    if "custom_skip_auto_creation" in existing_cols and "custom_skip_auto_creation" not in parsed:
        parsed["custom_skip_auto_creation"] = 0

    # Monthly Salary / CTC
    monthly_sal = row_dict.get("monthly_gross_salary") or row_dict.get("ctc") or row_dict.get("salary") or row_dict.get("monthly_salary")
    if monthly_sal and "ctc" in existing_cols and not parsed.get("ctc"):
        try:
            parsed["ctc"] = flt(monthly_sal)
        except Exception:
            pass

    # Status & Demographics normalization
    parsed["status"] = _normalize_status(parsed.get("status"), parsed.get("relieving_date"))
    if parsed.get("gender"):
        parsed["gender"] = _normalize_gender(parsed["gender"])
    if parsed.get("marital_status"):
        parsed["marital_status"] = _normalize_marital_status(parsed["marital_status"])

    # Company
    if not parsed.get("company"):
        parsed["company"] = frappe.defaults.get_global_default("company") or frappe.db.get_value("Company", {}, "name")

    # Filter out columns that don't exist in tabEmployee
    keys_to_remove = [k for k in list(parsed.keys()) if k != "doctype" and k not in existing_cols]
    for k in keys_to_remove:
        parsed.pop(k, None)

    # Create and Insert Doc
    doc = frappe.get_doc(parsed)
    doc.flags.ignore_links = True
    doc.flags.ignore_permissions = True
    doc.flags.ignore_mandatory = True
    doc.flags.ignore_version = True
    doc.insert(ignore_permissions=True, ignore_mandatory=True)

    # Auto Create User if not handled by after_insert hook
    if not doc.user_id and not doc.get("custom_skip_auto_creation"):
        try:
            from sahayog.doc_events.create_user_from_employee import create_user
            create_user(doc)
        except Exception:
            pass

    return doc.name


def _update_employee(emp_name, row_dict, field_map, cache=None, existing_cols=None):
    if existing_cols is None:
        existing_cols = set(r[0] for r in frappe.db.sql("SHOW COLUMNS FROM `tabEmployee`"))
    else:
        existing_cols = set(existing_cols)

    header_for = {v: k for k, v in field_map.items()}
    updates = {}

    date_fields = {
        "date_of_birth": _parse_date,
        "date_of_joining": _parse_date,
        "final_confirmation_date": _parse_date,
        "relieving_date": _parse_date,
        "resignation_letter_date": _parse_date,
    }

    link_map = {
        "custom_division": ("Division", "division", None),
        "custom_zone": ("Zone", "zone", "ZONE-"),
        "custom_region": ("Region", "region", "REGION-"),
        "branch": ("Branch", "branch", None),
        "designation": ("Designation", "designation_name", None),
    }

    skip_names = {"first_name", "middle_name", "last_name"}

    for csv_key, doc_field in field_map.items():
        csv_val = row_dict.get(csv_key)
        if not csv_val or doc_field in skip_names:
            continue

        if doc_field in date_fields:
            parsed = date_fields[doc_field](csv_val)
            if parsed and doc_field in existing_cols:
                updates[doc_field] = parsed
        elif doc_field == "ctc":
            try:
                if "ctc" in existing_cols:
                    updates["ctc"] = flt(csv_val)
            except Exception:
                pass
        elif doc_field in link_map:
            doctype, label, prefix = link_map[doc_field]
            link_res = _ensure_link(csv_val, doctype, label, prefix, cache=cache)
            if link_res and doc_field in existing_cols:
                updates[doc_field] = link_res
        elif doc_field == "sahayog_branch":
            clean_sb = csv_val.replace(" ", "")
            branch_info = cache["Sahayog Branch"].get(clean_sb) if cache else None
            if not branch_info:
                branch = frappe.db.sql(
                    """SELECT name FROM `tabSahayog Branch`
                       WHERE REPLACE(name, ' ', '') = %s LIMIT 1""",
                    clean_sb,
                )
                if branch:
                    branch_info = frappe.db.get_value(
                        "Sahayog Branch", branch[0][0],
                        ["name", "zone", "region", "district"], as_dict=True
                    )
                    if cache and "Sahayog Branch" in cache:
                        cache["Sahayog Branch"][clean_sb] = branch_info
            if branch_info and "sahayog_branch" in existing_cols:
                updates["sahayog_branch"] = branch_info.get("name") if isinstance(branch_info, dict) else branch_info
        elif doc_field == "department":
            dep_res = _ensure_department(csv_val, cache=cache)
            if dep_res and "department" in existing_cols:
                updates["department"] = dep_res
        elif doc_field == "reports_to":
            emp_lead = cache["Employee"].get(csv_val) if cache else None
            if not emp_lead:
                emp_lead = csv_val if frappe.db.exists("Employee", csv_val) else frappe.db.get_value("Employee", {"employee_number": csv_val}, "name")
                if emp_lead and cache:
                    cache["Employee"][csv_val] = emp_lead
            if emp_lead and "reports_to" in existing_cols:
                updates["reports_to"] = emp_lead
        elif doc_field == "status":
            relieving = updates.get("relieving_date") or row_dict.get("relieving_date")
            updates["status"] = _normalize_status(csv_val, relieving)
        elif doc_field == "gender":
            norm_g = _normalize_gender(csv_val)
            if norm_g and "gender" in existing_cols:
                updates["gender"] = norm_g
        elif doc_field == "marital_status":
            norm_m = _normalize_marital_status(csv_val)
            if norm_m and "marital_status" in existing_cols:
                updates["marital_status"] = norm_m
        else:
            if doc_field in existing_cols:
                updates[doc_field] = csv_val

    if "first_name" in field_map:
        csv_first_name = row_dict.get(header_for.get("first_name"))
        csv_middle_name = row_dict.get(header_for.get("middle_name"))
        csv_last_name = row_dict.get(header_for.get("last_name"))
        if csv_first_name:
            fn, mn, ln = _split_name(csv_first_name, csv_middle_name, csv_last_name)
            if fn and "first_name" in existing_cols:
                updates["first_name"] = fn
            if mn and "middle_name" in existing_cols:
                updates["middle_name"] = mn
            if ln and "last_name" in existing_cols:
                updates["last_name"] = ln

    sol_id = row_dict.get("sol_id")
    if sol_id and "sahayog_branch" in existing_cols and "sahayog_branch" not in updates:
        clean_sol = str(sol_id).strip().replace(" ", "")
        branch_info = cache.get("Sahayog Branch", {}).get(clean_sol) if cache else None
        if not branch_info:
            branch = frappe.db.sql(
                """SELECT name FROM `tabSahayog Branch` WHERE REPLACE(name, ' ', '') = %s LIMIT 1""",
                clean_sol,
            )
            if branch:
                branch_info = branch[0][0]
                if cache and "Sahayog Branch" in cache:
                    cache["Sahayog Branch"][clean_sol] = branch_info
        if branch_info:
            updates["sahayog_branch"] = branch_info.get("name") if isinstance(branch_info, dict) else branch_info

    relieving = updates.get("relieving_date")
    if relieving and getdate(relieving) <= getdate(today()):
        is_excluded = frappe.db.get_value("Employee", emp_name, "exclude_zinghr")
        if not is_excluded:
            updates["status"] = "Left"
            user_id = frappe.db.get_value("Employee", emp_name, "user_id")
            if user_id:
                frappe.db.set_value("User", user_id, "enabled", 0, update_modified=False)

    if updates:
        frappe.db.set_value("Employee", emp_name, updates, update_modified=True)


def _set_sol_fields(doc, row_dict, cache=None, existing_cols=None):
    if existing_cols is None:
        existing_cols = set(r[0] for r in frappe.db.sql("SHOW COLUMNS FROM `tabEmployee`"))
    else:
        existing_cols = set(existing_cols)

    sol_id = row_dict.get("sol_id")
    if sol_id:
        clean_sol = str(sol_id).strip().replace(" ", "")
        branch_info = None

        if cache and "Sahayog Branch" in cache:
            branch_info = cache["Sahayog Branch"].get(clean_sol)

        if not branch_info:
            branch = frappe.db.sql(
                """SELECT name FROM `tabSahayog Branch`
                   WHERE REPLACE(name, ' ', '') = %s LIMIT 1""",
                clean_sol,
            )
            if branch:
                branch_info = branch[0][0]
                if cache and "Sahayog Branch" in cache:
                    cache["Sahayog Branch"][clean_sol] = branch_info

        if branch_info:
            b_name = branch_info.get("name") if isinstance(branch_info, dict) else branch_info
            if "sahayog_branch" in existing_cols:
                doc.sahayog_branch = b_name

    monthly_sal = row_dict.get("monthly_gross_salary") or row_dict.get("ctc")
    if monthly_sal:
        try:
            if "ctc" in existing_cols and not doc.get("ctc"):
                doc.ctc = flt(monthly_sal)
        except Exception:
            pass


def _prepare_employee_data(row_dict, field_map):
    data = {"doctype": "Employee"}

    date_fields = {
        "date_of_birth": "date_of_birth",
        "date_of_joining": "date_of_joining",
        "final_confirmation_date": "final_confirmation_date",
        "relieving_date": "relieving_date",
        "resignation_letter_date": "resignation_letter_date",
    }

    for csv_key, doc_field in field_map.items():
        val = row_dict.get(csv_key)
        if not val:
            continue

        if doc_field in date_fields:
            parsed = _parse_date(val)
            if parsed:
                data[doc_field] = parsed
        elif doc_field == "cell_number":
            data["cell_number"] = str(val)
        elif doc_field == "status":
            data["status"] = _normalize_status(val, data.get("relieving_date"))
        elif doc_field == "gender":
            data["gender"] = _normalize_gender(val)
        elif doc_field == "marital_status":
            data["marital_status"] = _normalize_marital_status(val)
        elif doc_field == "ctc":
            try:
                data["ctc"] = flt(val)
            except Exception:
                pass
        else:
            data[doc_field] = val

    if "status" not in data:
        data["status"] = _normalize_status(None, data.get("relieving_date"))

    return data


def _normalize_status(val, relieving_date=None):
    if relieving_date:
        try:
            rd = getdate(relieving_date) if isinstance(relieving_date, str) else relieving_date
            if rd and rd <= getdate(today()):
                return "Left"
        except Exception:
            pass

    if not val:
        return "Active"

    s = str(val).strip().lower()
    if s in ("active", "newjoinee", "new joinee", "new_joinee", "existing", "confirmed", "probation", "probationary", "joined", "onboarded", "regular"):
        return "Active"
    elif s in ("left", "resigned", "terminated", "relieved", "absconded", "retired"):
        return "Left"
    elif s in ("inactive", "on leave", "disabled"):
        return "Inactive"
    elif s in ("suspended",):
        return "Suspended"
    return "Active"


def _normalize_gender(val):
    if not val:
        return None
    g = str(val).strip().lower()
    if g in ("m", "male"):
        return "Male"
    elif g in ("f", "female"):
        return "Female"
    elif g in ("transgender", "trans"):
        return "Transgender"
    elif g in ("other",):
        return "Other"
    return val.strip().title()


def _normalize_marital_status(val):
    if not val:
        return None
    m = str(val).strip().lower()
    if m in ("single", "unmarried"):
        return "Single"
    elif m in ("married",):
        return "Married"
    elif m in ("divorced",):
        return "Divorced"
    elif m in ("widowed", "widow"):
        return "Widowed"
    return val.strip().title()


def _parse_date(val):
    if val is None:
        return None

    if isinstance(val, (datetime.date, datetime.datetime)):
        return val.strftime("%Y-%m-%d")

    val_str = str(val).strip()
    if not val_str or val_str.lower() in ("none", "null", "na", "n/a", "-", "", "nan"):
        return None

    # Strip any time component
    if " " in val_str:
        val_str = val_str.split(" ")[0].strip()
    elif "t" in val_str.lower():
        val_str = val_str.lower().split("t")[0].strip()

    # Handle Excel float/int serial dates (e.g. 44562 or 44562.0)
    try:
        f_val = float(val_str)
        if 20000 <= f_val <= 65000:
            excel_base = datetime.date(1899, 12, 30)
            parsed_date = excel_base + datetime.timedelta(days=int(f_val))
            return parsed_date.strftime("%Y-%m-%d")
    except (ValueError, OverflowError):
        pass

    # Try explicit standard formats first
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%Y/%m/%d", "%d-%b-%Y", "%d-%B-%Y"):
        try:
            return datetime.datetime.strptime(val_str, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue

    # Fallback to Frappe getdate parser
    try:
        d = getdate(val_str)
        if d:
            return get_datetime_str(d).split(" ")[0]
    except Exception:
        pass

    return None


def _build_summary(result, mode):
    now = now_datetime()
    date_str = now.strftime("%d/%m/%Y %I:%M %p")

    lines = []
    lines.append(f"Import on {date_str}")
    lines.append("")

    if mode == "insert":
        lines.append(f"Inserted: {result['inserted']}")
        if result["inserted"]:
            nums = sorted(result["inserted_numbers"])
            if len(nums) == 1:
                lines.append(f"  Employee No: {nums[0]}")
            elif nums:
                lines.append(f"  Employee No: {nums[0]} to {nums[-1]}")
    else:
        lines.append(f"Updated: {result['updated']}")
        if result["updated"]:
            nums = sorted(result["updated_numbers"])
            if len(nums) == 1:
                lines.append(f"  Employee No: {nums[0]}")
            elif nums:
                lines.append(f"  Employee No: {nums[0]} to {nums[-1]}")

    lines.append(f"Skipped: {result['skipped']}")
    lines.append(f"Failed: {result['failed']}")

    if result["errors"]:
        lines.append("")
        lines.append("Errors:")
        for err in result["errors"][:20]:
            lines.append(f"  {err}")
        if len(result["errors"]) > 20:
            lines.append(f"  ... and {len(result['errors']) - 20} more")

    return "\n".join(lines)


@frappe.whitelist()
def get_file_headers():
    """Read headers from uploaded Employee Master file and return smart target field suggestions."""
    setting = frappe.get_doc("Sahayog HR Setting")
    file_url = setting.get("employee_master")
    if not file_url:
        return {"headers": [], "mappings": []}

    rows = _parse_file(file_url)
    if not rows or len(rows) < 1:
        return {"headers": [], "mappings": []}

    headers = [str(h).strip() for h in rows[0] if h and str(h).strip()]
    existing_cols = set(r[0] for r in frappe.db.sql("SHOW COLUMNS FROM `tabEmployee`"))

    mappings = []
    for h in headers:
        clean = _clean_header(h)
        target = FIELD_MAP.get(clean, clean)
        if target not in existing_cols and f"custom_{target}" in existing_cols:
            target = f"custom_{target}"
        elif target not in existing_cols:
            target = ""
        mappings.append({
            "source_column": h,
            "target_field": target,
        })

    return {"headers": headers, "mappings": mappings}


class _StopRow(Exception):
    pass
