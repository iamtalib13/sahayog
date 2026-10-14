import io
import json
from html import escape

import frappe
from sahayog.branch_score_card.audit_log_utils import get_deleted_record_audit_logs, audit_details, audit_details_html, audit_details_text

DT = "Account Opening Operations"
TABLES = {
    "table_dllf": "FTR / FTNR",
    "table_zero_ip_funding": "Zero IP Funding",
}
SKIP = {
    "name", "owner", "creation", "modified", "modified_by", "docstatus",
    "parent", "parentfield", "parenttype", "idx", "doctype", "__unsaved",
    "scheme_code_options",
}
EMPTY = ("(empty)", "0", "0.0")

TEMPLATE = """<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Account Opening Operations - Logs</title>
<style>
*{box-sizing:border-box}
body{margin:0;background:#fff;color:#1f2933;font:13px/1.5 -apple-system,"Segoe UI",Roboto,Arial,sans-serif}
.wrap{width:100%;padding:24px 32px}
.top{display:flex;justify-content:space-between;align-items:flex-end;gap:16px;flex-wrap:wrap;margin:0 0 14px}
.top h1{margin:0;font-size:18px;font-weight:600}
.sub{color:#8a94a0;font-size:12px;margin-top:2px}
.meta{display:flex;gap:18px;align-items:center;flex-wrap:wrap;color:#8a94a0;font-size:12px}
.pill b{color:#1f2933;font-weight:600;margin-right:4px}
#q{border:0;border-bottom:1px solid #d5d9de;padding:4px 2px;width:220px;font-size:13px;background:none;outline:none}
#q:focus{border-color:#1f2933}
.dl_bar{display:flex;gap:18px;align-items:center;flex-wrap:wrap;margin:0 0 6px;font-size:12px;color:#8a94a0}
.dl_bar input{border:0;border-bottom:1px solid #d5d9de;padding:3px 2px;font-size:13px;background:none;color:#1f2933;outline:none}
.dl_bar input:focus{border-color:#1f2933}
.dl_bar button{border:1px solid #1f2933;background:none;color:#1f2933;border-radius:4px;padding:4px 12px;font-size:12px;cursor:pointer}
.dl_bar button:hover{background:#1f2933;color:#fff}
table{width:100%;border-collapse:collapse}
thead th{text-align:left;font-size:11px;font-weight:500;text-transform:uppercase;letter-spacing:.06em;color:#8a94a0;padding:10px 12px;border-bottom:1px solid #e5e7eb}
tbody td{padding:10px 12px;border-bottom:1px solid #f0f1f3;vertical-align:top}
tbody tr:last-child td{border-bottom:0}
.when{white-space:nowrap;width:140px;color:#6b7785}
.rec{width:190px;font-weight:500}
.who{width:150px}
.who b{display:block;font-weight:500}
.who span{color:#a0a8b1;font-size:12px}
.act{width:90px}
.badge{font-size:12px;font-weight:500;color:#52606d}
.badge:before{content:"";display:inline-block;width:6px;height:6px;border-radius:50%;margin-right:6px;vertical-align:middle;background:#a0a8b1}
.badge.created{color:#067647}.badge.created:before{background:#067647}
.badge.updated{color:#1d56b8}.badge.updated:before{background:#1d56b8}
.chgs{display:flex;flex-wrap:wrap;gap:2px 22px}
.tb{margin-right:6px;color:#8a94a0;font-size:12px}
.f{font-weight:500}
.old{color:#b42318;text-decoration:line-through}
.new{color:#067647;font-weight:500}
.arr{color:#a0a8b1;margin:0 4px}
.add{color:#067647;font-weight:500}
.rem{color:#b42318;font-weight:500}
.muted{color:#a0a8b1}
.empty{padding:28px;text-align:center;color:#a0a8b1;display:none}
</style></head><body><div class="wrap">
<div class="top">
<div><h1>Account Opening Operations - Activity Logs</h1><div class="sub">All branches &middot; latest first</div></div>
<div class="meta">
<input id="q" type="text" placeholder="Search SOL / month / user...">
<span class="pill"><b>__N__</b>Events</span>
<span class="pill"><b>__B__</b>Records</span>
<span class="pill"><b>__U__</b>Users</span>
</div></div>
<div class="dl_bar" id="dl_bar" data-url="/api/method/sahayog.branch_score_card.doctype.account_opening_operations.aoo_logs_page.download_logs">
<label>From <input type="date" id="dl_from"></label>
<label>To <input type="date" id="dl_to"></label>
<button type="button" id="dl_btn">Download Excel</button></div>
<table>
<thead><tr><th>Date / Time</th><th>Record</th><th>User</th><th>Action</th><th>What changed</th></tr></thead>
<tbody id="tb">__ROWS__</tbody>
</table><div class="empty" id="empty">No matching logs</div>
</div>
<script>
var q=document.getElementById('q'),rows=document.querySelectorAll('#tb tr'),em=document.getElementById('empty');
q.addEventListener('input',function(){var v=q.value.toLowerCase().trim(),n=0;
rows.forEach(function(r){var ok=!v||r.getAttribute('data-s').indexOf(v)>-1;r.style.display=ok?'':'none';if(ok)n++;});
em.style.display=n?'none':'block';});
(function(){var b=document.getElementById('dl_bar');
document.getElementById('dl_btn').onclick=function(){
var f=document.getElementById('dl_from').value,t=document.getElementById('dl_to').value;
if(f&&t&&f>t){alert('From Date cannot be after To Date');return;}
var p=[];if(f)p.push('from_date='+f);if(t)p.push('to_date='+t);
window.location=b.dataset.url+(p.length?'?'+p.join('&'):'');};})();
</script></body></html>"""


def _can_view():
    u = frappe.session.user
    return u == "Administrator" or bool(
        {"System Manager", "Branch Scorecard Editor"} & set(frappe.get_roles(u))
    )


def _v(x):
    return "(empty)" if x is None or x == "" else str(x)


def _lm(doctype):
    return {f.fieldname: (f.label or f.fieldname) for f in frappe.get_meta(doctype).fields}


def _collect(from_date=None, to_date=None, limit=0):
    """Flat log items: (time, record, user, action, table, row, field, old, new, details)"""
    from frappe.utils import add_days, getdate

    fd = getdate(from_date) if from_date else None
    td = getdate(to_date) if to_date else None
    if fd and td and fd > td:
        frappe.throw(frappe._("From Date cannot be after To Date"))

    date_f = []
    if fd:
        date_f.append(["creation", ">=", fd])
    if td:
        date_f.append(["creation", "<", add_days(td, 1)])

    parent_l = _lm(DT)
    child_dt = {f.fieldname: f.options for f in frappe.get_meta(DT).fields if f.fieldname in TABLES}
    child_l = {t: _lm(o) for t, o in child_dt.items()}

    def pairs(t, row):
        return ", ".join(
            "%s: %s" % (child_l.get(t, {}).get(k, k), v)
            for k, v in row.items()
            if k not in SKIP and v not in (None, "") and str(v) not in ("0", "0.0")
        ) or "blank row"

    def sig(row):
        return tuple(sorted(
            (k, str(v)) for k, v in row.items() if k not in SKIP and v not in (None, "")
        ))

    items = []
    for d in frappe.get_all(DT, filters=date_f, fields=["name", "owner", "creation"], limit_page_length=0):
        items.append((d.creation, d.name, d.owner, "Created", "", "", "", "", "", "Record created"))

    versions = frappe.get_all(
        "Version",
        filters=[["ref_doctype", "=", DT]] + date_f,
        fields=["creation", "owner", "docname", "data"],
        order_by="creation desc",
        limit_page_length=limit or 0,
    )
    for v in versions:
        try:
            data = json.loads(v.data or "{}")
        except Exception:
            continue
        deleted_snapshot = data.get("deleted")
        if isinstance(deleted_snapshot, dict):
            items.append((
                v.creation, v.docname, v.owner, "Deleted", "Record", "",
                "Deleted snapshot", "", "",
                json.dumps(deleted_snapshot, ensure_ascii=False, default=str),
            ))
            continue

        base = (v.creation, v.docname, v.owner, "Updated")

        for ch in data.get("changed") or []:
            if len(ch) >= 3 and not (_v(ch[1]) in EMPTY and _v(ch[2]) in EMPTY):
                items.append(base + ("Record", "", parent_l.get(ch[0], ch[0]), _v(ch[1]), _v(ch[2]), ""))

        # validate() rows dobara banata hai -> identical added/removed pair ko noise maan ke hatao
        added = [(i[0], i[1]) for i in (data.get("added") or []) if len(i) >= 2 and isinstance(i[1], dict)]
        removed = [(i[0], i[1]) for i in (data.get("removed") or []) if len(i) >= 2 and isinstance(i[1], dict)]
        rem_pool = [(t, sig(r)) for t, r in removed]
        real_added = []
        for t, r in added:
            key = (t, sig(r))
            if key in rem_pool:
                rem_pool.remove(key)
            else:
                real_added.append((t, r))
        real_removed = [(t, r) for t, r in removed if (t, sig(r)) in rem_pool]

        for t, r in real_added:
            items.append(base + (TABLES.get(t, t), "", "Row Added", "", "", pairs(t, r)))
        for t, r in real_removed:
            items.append(base + (TABLES.get(t, t), "", "Row Removed", "", "", pairs(t, r)))

        for item in data.get("row_changed") or []:
            if len(item) < 4:
                continue
            t, idx, changes = item[0], item[1], item[3]
            for c in changes or []:
                if c[0] in SKIP:
                    continue
                if len(c) >= 3 and not (_v(c[1]) in EMPTY and _v(c[2]) in EMPTY):
                    items.append(base + (
                        TABLES.get(t, t), int(idx) + 1,
                        child_l.get(t, {}).get(c[0], c[0]), _v(c[1]), _v(c[2]), "",
                    ))

    permanent_rows = get_deleted_record_audit_logs(DT)
    deleted_names = {row.record_name for row in permanent_rows}

    # Preserve existing child diffs for live records; replace deleted-record events.
    items = [item for item in items if item[1] not in deleted_names]

    for audit in permanent_rows:
        if fd and frappe.utils.getdate(audit.event_time) < fd:
            continue
        if td and frappe.utils.getdate(audit.event_time) > td:
            continue
        items.append((
            audit.event_time,
            audit.record_name,
            audit.user,
            audit.action,
            "Record",
            "",
            "Permanent audit details",
            "",
            "",
            audit_details_text(audit),
        ))

    items.sort(key=lambda x: x[0], reverse=True)
    return items


@frappe.whitelist()
def view_logs_page():
    if not _can_view():
        frappe.throw(frappe._("Not permitted"), frappe.PermissionError)

    items = _collect(limit=3000)

    groups = {}
    for it in items:
        groups.setdefault((it[0], it[1], it[2], it[3]), []).append(it)

    names = {}

    def fullname(u):
        if u not in names:
            fn = frappe.utils.get_fullname(u)
            names[u] = fn if fn and fn != u else ""
        return names[u]

    audit_map = {
        (a.event_time, a.record_name, a.user, a.action): a
        for a in get_deleted_record_audit_logs(DT)
    }

    rows = []
    for (time, record, user, action), lst in groups.items():
        perm_audit = audit_map.get((time, record, user, action))
        if perm_audit is not None and not any(
            it[6] == "Permanent audit details" for it in lst
        ):
            perm_audit = None
        if perm_audit is not None:
            html = (
                "<div class='chgs'><div class='ln'>"
                "<span class='tb'>Permanent audit \u00b7 Source: %s</span>"
                "%s</div></div>"
            ) % (escape(perm_audit.source or "Manual"), audit_details_html(perm_audit))
        elif action == "Created":
            html = "<span class='muted'>Record created</span>"
        elif action == "Deleted":
            details = next((it[9] for it in lst if it[9]), "")
            html = (
                "<div class='chgs'><div class='ln'>"
                "<span class='tb'>Deleted snapshot</span>"
                "<span class='f'>%s</span></div></div>" % escape(details)
            )
        else:
            parts = []
            for it in lst:
                tb = escape(str(it[4]))
                if it[5] != "":
                    tb += " &middot; Row %s" % it[5]
                if it[6] in ("Row Added", "Row Removed"):
                    cls = "add" if it[6] == "Row Added" else "rem"
                    parts.append(
                        "<div class='ln'><span class='tb'>%s</span><span class='%s'>%s</span> &middot; %s</div>"
                        % (tb, cls, it[6], escape(it[9]))
                    )
                else:
                    parts.append(
                        "<div class='ln'><span class='tb'>%s</span><span class='f'>%s</span>: "
                        "<span class='old'>%s</span><span class='arr'>&rarr;</span><span class='new'>%s</span></div>"
                        % (tb, escape(str(it[6])), escape(it[7]), escape(it[8]))
                    )
            html = "<div class='chgs'>" + "".join(parts) + "</div>"

        fn = fullname(user)
        when = frappe.utils.format_datetime(time, "dd-MM-yyyy HH:mm:ss")
        search = escape(("%s %s %s" % (record, user, fn)).lower(), quote=True)
        rows.append(
            "<tr data-s=\"%s\"><td class='when'>%s</td><td class='rec'>%s</td>"
            "<td class='who'><b>%s</b><span>%s</span></td>"
            "<td class='act'><span class='badge %s'>%s</span></td><td>%s</td></tr>"
            % (search, escape(when), escape(record), escape(fn or user or ""),
               escape(user or ""), action.lower(), action, html)
        )

    page = (
        TEMPLATE.replace("__N__", str(len(groups)))
        .replace("__B__", str(len({k[1] for k in groups})))
        .replace("__U__", str(len({k[2] for k in groups})))
        .replace("__ROWS__", "".join(rows))
    )

    frappe.local.response.filename = "account_opening_logs.html"
    frappe.local.response.filecontent = page
    frappe.local.response.type = "download"
    frappe.local.response.display_content_as = "inline"


@frappe.whitelist()
def download_logs(from_date=None, to_date=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    if not _can_view():
        frappe.throw(frappe._("Not permitted"), frappe.PermissionError)

    items = _collect(from_date, to_date)
    names = {}

    def full(u):
        if u not in names:
            fn = frappe.utils.get_fullname(u)
            names[u] = fn if fn and fn != u else ""
        return names[u]

    wb = Workbook()
    ws = wb.active
    ws.title = "Logs"
    ws.append(["Date / Time", "Record", "User", "Name", "Action", "Table", "Row", "Field", "Old Value", "New Value", "Details"])
    for it in items:
        ws.append([
            frappe.utils.format_datetime(it[0], "dd-MM-yyyy HH:mm:ss"), it[1], it[2], full(it[2]), it[3],
            it[4], it[5], it[6], it[7], it[8], it[9],
        ])

    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        w = max(len(str(c.value)) if c.value is not None else 0 for c in col[:300])
        ws.column_dimensions[col[0].column_letter].width = min(max(w + 2, 12), 60)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)

    buf = io.BytesIO()
    wb.save(buf)
    frappe.response["filename"] = "Account_Opening_Logs_%s.xlsx" % frappe.utils.today()
    frappe.response["filecontent"] = buf.getvalue()
    frappe.response["type"] = "binary"
