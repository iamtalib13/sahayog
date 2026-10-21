# Copyright (c) 2026, Developer Team and contributors
# For license information, please see license.txt

import frappe
from frappe.model.document import Document
from frappe.utils import flt, getdate, today


def get_current_fy():
    """Financial year (Apr-Mar) in YYYY-YY format, e.g. 2026-27"""
    d = getdate(today())
    start = d.year if d.month >= 4 else d.year - 1
    return f"{start}-{str(start + 1)[-2:]}"


class COMVisitTracker(Document):
    def before_insert(self):
        if not self.year:
            self.year = get_current_fy()

    def validate(self):
        for row in self.get("com_visit_tracker") or []:
            month = (row.month or "").strip()

            if month:
                for fieldname, label in (
                    ("visti_start_date", "Visit Start Date"),
                    ("report_publish_date", "Report Publish Date"),
                ):
                    date_value = row.get(fieldname)
                    if date_value and getdate(date_value).strftime("%B").lower() != month.lower():
                        actual_month = getdate(date_value).strftime("%B")
                        frappe.throw(
                            frappe._(
                                "Row {0}: Month is {1}, but {2} is in {3}. "
                                "Please correct the date to match the selected month before saving."
                            ).format(row.idx, month, label, actual_month),
                            title=frappe._("Month Mismatch"),
                        )

            if (
                row.report_publish_date
                and row.report_closure_date
                and getdate(row.report_closure_date) < getdate(row.report_publish_date)
            ):
                frappe.throw(
                    frappe._("Row {0}: Report Closure Date cannot be earlier than Report Publish Date.").format(row.idx)
                )

    def on_update(self):
        try:
            sync_to_audit_and_compliance(self)
        except Exception:
            frappe.log_error(
                title=f"COM Visit Tracker sync failed [{self.name}]",
                message=frappe.get_traceback(),
            )

    def autoname(self):
        if not self.year:
            self.year = get_current_fy()
        b = frappe.db.get_value(
            "Sahayog Branch", self.sol_id,
            ["sol_id", "branch", "district"], as_dict=True,
        ) or {}
        parts = [
            str(b.get("sol_id") or self.sol_id),
            b.get("branch") or "",
            b.get("district") or "",
            self.year,
        ]
        self.name = "-".join(p.replace("/", "-").strip() for p in parts)


def create_yearly_trackers():
    """Naye saal ke liye: sirf un SOLs ke records jinka kisi COM ke Report Preference me access hai."""
    created = 0
    for name in frappe.get_all("Report Preference", filters={"enabled": 1}, pluck="name"):
        ensure_trackers_for_preference(frappe.get_doc("Report Preference", name))
    frappe.db.commit()
    return created


# Jin designations ko access milega (lowercase me likho)
COM_DESIGNATIONS = {"cluster operation manager"}


def is_com(user):
    """Sirf Cluster Operation Manager designation wale user ko access."""
    d = frappe.db.get_value("Employee", {"user_id": user}, "designation")
    return (d or "").strip().lower() in COM_DESIGNATIONS


def _is_admin(user):
    return user == "Administrator" or bool(
        {"System Manager", "Branch Scorecard Editor"} & set(frappe.get_roles(user))
    )


def _get_scope(user):
    """Report Preference se zone/region/district/sol_id. None = koi access nahi."""
    if not is_com(user):
        return None
    if not frappe.db.exists("Report Preference", user):
        return None
    pref = frappe.get_doc("Report Preference", user)
    if not pref.enabled:
        return None

    def vals(table, field):
        return [str(r.get(field)) for r in (pref.get(table) or []) if r.get(field)]

    scope = {
        "geo": {
            "zone": vals("zone", "zone"),
            "region": vals("region", "region"),
            "state": vals("state", "state"),
            "district": vals("district", "district"),
        },
        "sol_ids": vals("sol_id", "sol_id"),
    }
    if pref.access_type == "Specific Branches (SOL ID)":
        scope["geo"] = {k: [] for k in scope["geo"]}
    else:
        scope["sol_ids"] = []
    return scope


def get_permission_query_conditions(user=None):
    user = user or frappe.session.user
    if _is_admin(user):
        return ""
    scope = _get_scope(user)
    if not scope:
        return "1=0"

    t = "`tabCOM Visit Tracker`"
    geo = [
        f"{t}.`{f}` in ({', '.join(frappe.db.escape(v) for v in vs)})"
        for f, vs in scope["geo"].items() if vs
    ]
    clauses = []
    if geo:
        clauses.append("(" + " and ".join(geo) + ")")
    if scope["sol_ids"]:
        in_sols = ", ".join(frappe.db.escape(v) for v in scope["sol_ids"])
        clauses.append(f"{t}.`sol_id` in ({in_sols})")
    return "(" + " or ".join(clauses) + ")" if clauses else "1=0"


def has_permission(doc, ptype=None, user=None, debug=False):
    user = user or frappe.session.user
    if _is_admin(user):
        return True
    scope = _get_scope(user)
    if not scope:
        return False
    if doc.get("sol_id") and str(doc.sol_id) in scope["sol_ids"]:
        return True
    geo = {f: vs for f, vs in scope["geo"].items() if vs}
    return bool(geo) and all(doc.get(f) in vs for f, vs in geo.items())


def get_scope_sols(pref):
    """Report Preference ke access se saare SOL (Zonal chhodke)."""
    base = [["branch_type", "!=", "Zonal"]]
    geo = {}
    for f in ("zone", "region", "state", "district"):
        vs = [r.get(f) for r in (pref.get(f) or []) if r.get(f)]
        if vs:
            geo[f] = vs
    sols = [str(r.sol_id) for r in (pref.get("sol_id") or []) if r.get("sol_id")]
    if pref.access_type == "Specific Branches (SOL ID)":
        geo = {}
    else:
        sols = []

    found = set()
    if geo:
        filters = base + [[f, "in", vs] for f, vs in geo.items()]
        found.update(frappe.get_all("Sahayog Branch", filters=filters, pluck="name"))
    if sols:
        filters = base + [["name", "in", sols]]
        found.update(frappe.get_all("Sahayog Branch", filters=filters, pluck="name"))
    return found


def ensure_trackers_for_preference(doc, method=None):
    """Report Preference save hote hi us user ke SOLs ke records bana do (jo nahi hain)."""
    if not doc.enabled or not is_com(doc.user):
        return
    year = get_current_fy()
    sols = get_scope_sols(doc)
    if not sols:
        return
    existing = set(frappe.get_all(
        "COM Visit Tracker",
        filters={"year": year, "sol_id": ["in", list(sols)]},
        pluck="sol_id",
    ))
    for sol in sols - existing:
        try:
            frappe.get_doc({
                "doctype": "COM Visit Tracker",
                "sol_id": sol,
                "year": year,
            }).insert(ignore_permissions=True)
        except Exception:
            frappe.log_error(
                frappe.get_traceback(),
                f"COM Visit Tracker create failed: {sol}",
            )


def backfill_all_preferences():
    """Jo Report Preferences pehle se hain unke liye ek baar chalao."""
    for name in frappe.get_all("Report Preference", pluck="name"):
        ensure_trackers_for_preference(frappe.get_doc("Report Preference", name))
    frappe.db.commit()


def create_trackers_for_all_branches(dry_run=True):
    """Saare Sahayog Branch (Zonal chhodke) ke current FY ke tracker banata hai."""
    year = get_current_fy()
    all_sols = set(frappe.get_all(
        "Sahayog Branch", filters={"branch_type": ["!=", "Zonal"]}, pluck="name"
    ))
    existing = set(frappe.get_all(
        "COM Visit Tracker", filters={"year": year}, pluck="sol_id"
    ))
    missing = sorted(all_sols - existing)
    if dry_run:
        return len(missing)
    for sol in missing:
        try:
            frappe.get_doc({
                "doctype": "COM Visit Tracker",
                "sol_id": sol,
                "year": year,
            }).insert(ignore_permissions=True)
        except Exception:
            frappe.log_error(frappe.get_traceback(), f"COM Visit Tracker create failed: {sol}")
    frappe.db.commit()
    return len(missing)


# ---------------------------------------------------------------
# COM Visit Tracker -> Audit and Compliance (one way, no delete)
# ---------------------------------------------------------------
def _mkey(v):
    return (v or "").strip().lower()


def _find_month_row(rows, month):
    key = _mkey(month)
    for r in rows:
        if _mkey(r.month) == key:
            return r
    return None


def _apply(ac, table, month, values):
    """Sirf bhari hui values likho. True return karta hai agar kuch badla."""
    if not values:
        return False
    row = _find_month_row(ac.get(table) or [], month)
    if not row:
        ac.append(table, {"month": month, **values})
        return True
    changed = False
    for k, v in values.items():
        if str(row.get(k) or "") != str(v):
            row.set(k, v)
            changed = True
    return changed


_MONTHS = ["january", "february", "march", "april", "may", "june", "july",
           "august", "september", "october", "november", "december"]


def _row_year(date_val, month, fy_start):
    """A&C year = date ka calendar year. Date na ho to month se (Jan-Mar = FY start + 1)."""
    if date_val:
        return getdate(date_val).year
    m = (month or "").strip().lower()
    if m in _MONTHS:
        return fy_start + 1 if _MONTHS.index(m) < 3 else fy_start
    return fy_start


def sync_to_audit_and_compliance(doc):
    if not doc.sol_id or not doc.year:
        return
    try:
        fy_start = int(str(doc.year).split("-")[0].strip())
    except ValueError:
        frappe.throw(frappe._("Invalid year format: {0}").format(doc.year))

    work = {}
    for r in doc.get("com_visit_tracker") or []:
        month = (r.month or "").strip()
        if not month:
            continue
        visit, comp = {}, {}
        if r.visti_start_date:
            visit["date_of_visit"] = str(getdate(r.visti_start_date))
        if flt(r.score_obtain):
            visit["visit_score_i"] = str(flt(r.score_obtain))
        if r.report_publish_date:
            comp["date_of_publish"] = str(getdate(r.report_publish_date))
        if r.report_closure_date:
            comp["date_of_closure"] = str(getdate(r.report_closure_date))
        if visit:
            y = _row_year(r.visti_start_date, month, fy_start)
            work.setdefault(y, []).append((month, visit, {}))
        if comp:
            y = _row_year(r.report_publish_date or r.report_closure_date, month, fy_start)
            work.setdefault(y, []).append((month, {}, comp))

    for ac_year, items in work.items():
        _sync_one_year(doc, ac_year, items)


def _sync_one_year(doc, ac_year, items):
    name = frappe.db.get_value(
        "Audit and Compliance", {"sol_id": doc.sol_id, "year": ac_year}, "name"
    )
    if name:
        ac = frappe.get_doc("Audit and Compliance", name)
    else:
        branch = frappe.db.get_value("Sahayog Branch", doc.sol_id, "branch") or doc.branch_name or "Unknown"
        ac = frappe.get_doc({
            "doctype": "Audit and Compliance",
            "sol_id": doc.sol_id,
            "branch_name": branch,
            "year": ac_year,
        })

    changed = False
    for month, visit, comp in items:
        changed |= _apply(ac, "com_visit", month, visit)
        changed |= _apply(ac, "com_visit_compliance", month, comp)

    if not changed:
        return
    if name:
        ac.save(ignore_permissions=True)
    else:
        ac.insert(ignore_permissions=True)
    frappe.msgprint(
        frappe._("Audit and Compliance updated: {0}").format(ac.name),
        alert=True, indicator="green",
    )


# ---------------------------------------------------------------
# Report Download: parent + child table ka poora data Excel mein
# ---------------------------------------------------------------
@frappe.whitelist()
def download_report(drafts=None):
    import io
    from decimal import Decimal
    from datetime import date, datetime, time, timedelta
    from frappe.model import no_value_fields
    from openpyxl import Workbook
    from openpyxl.styles import Font

    dt = "COM Visit Tracker"
    if not frappe.has_permission(dt, "read"):
        frappe.throw(frappe._("Not permitted"), frappe.PermissionError)

    meta = frappe.get_meta(dt)
    p_fields = [f for f in meta.fields if f.fieldtype not in no_value_fields]
    tables = [f for f in meta.fields if f.fieldtype == "Table"]

    # get_list = permission query conditions apply (admin/editor ko sab, COM ko apna scope)
    parents = frappe.get_list(
        dt,
        fields=["name"] + [f.fieldname for f in p_fields],
        order_by="sol_id asc",
        limit_page_length=0,
    )
    names = [p.name for p in parents]

    child_data = {}
    for t in tables:
        cmeta = frappe.get_meta(t.options)
        c_fields = [f for f in cmeta.fields if f.fieldtype not in no_value_fields]
        rows = frappe.get_all(
            t.options,
            filters={"parent": ["in", names or [""]], "parentfield": t.fieldname},
            fields=["parent", "idx"] + [f.fieldname for f in c_fields],
            order_by="parent asc, idx asc",
            limit_page_length=0,
        )
        by_parent = {}
        for r in rows:
            by_parent.setdefault(r.parent, []).append(r)
        child_data[t.fieldname] = (t, c_fields, by_parent)

    def clean(v):
        if v is None:
            return ""
        if isinstance(v, Decimal):
            return float(v)
        if isinstance(v, (time, timedelta)):
            return str(v)
        if isinstance(v, (datetime, date)):
            return v
        return v

    headers = ["ID", "Status"] + [f.label or f.fieldname for f in p_fields]
    for fn, (t, c_fields, _) in child_data.items():
        prefix = t.label or fn
        headers += [f.label or f.fieldname for f in c_fields]

    try:
        drafts_set = set(frappe.parse_json(drafts) or [])
    except Exception:
        drafts_set = set()

    def get_status(name):
        if name in drafts_set:
            return "Final Submission Pending"
        if any(name in bp for (_t, _c, bp) in child_data.values()):
            return "Submitted"
        return "Action Pending"

    wb = Workbook()
    ws = wb.active
    ws.title = "COM Visit Tracker"
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(name="Arial", size=9, bold=False)

    for p in parents:
        base = [clean(p.name), get_status(p.name)] + [clean(p.get(f.fieldname)) for f in p_fields]
        # har child table ke liye blank block (width fixed rakhne ke liye)
        blanks = {fn: [""] * len(cf) for fn, (t, cf, _) in child_data.items()}
        emitted = False
        for fn, (t, c_fields, by_parent) in child_data.items():
            for r in by_parent.get(p.name, []):
                block = [clean(r.get(f.fieldname)) for f in c_fields]
                row = list(base)
                for fn2 in child_data:
                    row += block if fn2 == fn else blanks[fn2]
                ws.append(row)
                emitted = True
        if not emitted:
            row = list(base)
            for fn2 in child_data:
                row += blanks[fn2]
            ws.append(row)

    from openpyxl.styles import PatternFill, Alignment, Border, Side

    thin = Side(style="thin", color="C9D1D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(name="Arial", bold=False, color="FFFFFF", size=9)

    # header colors: parent = dark blue, har child table ka alag color
    group_colors = ["2E7D32", "B45309", "6A1B9A", "00838F"]
    col_colors = ["1F4E78"] * (2 + len(p_fields))
    for i, (fn, (t, c_fields, _)) in enumerate(child_data.items()):
        col_colors += [group_colors[i % len(group_colors)]] * len(c_fields)

    for idx, c in enumerate(ws[1]):
        c.fill = PatternFill("solid", fgColor=col_colors[idx])
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    ws.row_dimensions[1].height = 34

    status_style = {
        "Submitted": ("C8E6C9", "1B5E20"),
        "Action Pending": ("FFCDD2", "B71C1C"),
        "Final Submission Pending": ("FFE0B2", "E65100"),
    }
    bands = ["EAF2FA", "FFFFFF"]
    band = 0
    last_id = None
    for r in range(2, ws.max_row + 1):
        rid = ws.cell(row=r, column=1).value
        if rid != last_id:
            band = 1 - band
            last_id = rid
        fill = PatternFill("solid", fgColor=bands[band])
        for c in ws[r]:
            c.fill = fill
            c.border = border
            c.font = Font(name="Arial", size=9, bold=False)
            c.alignment = Alignment(vertical="center", wrap_text=True)
            if isinstance(c.value, datetime):
                c.number_format = "DD-MM-YYYY HH:MM"
            elif isinstance(c.value, date):
                c.number_format = "DD-MM-YYYY"
        sc = ws.cell(row=r, column=2)
        bg, fg = status_style.get(sc.value, ("FFFFFF", "000000"))
        sc.fill = PatternFill("solid", fgColor=bg)
        sc.font = Font(name="Arial", size=9, bold=False, color=fg)
        sc.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=1).font = Font(name="Arial", size=9, bold=False)

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.zoomScale = 90
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:200])
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 14), 40)

    buf = io.BytesIO()
    wb.save(buf)
    frappe.response["filename"] = f"COM_Visit_Tracker_{today()}.xlsx"
    frappe.response["filecontent"] = buf.getvalue()
    frappe.response["type"] = "binary"

# ---------------------------------------------------------------
# View Logs (sirf System Manager / Administrator / Branch Scorecard Editor)
# ---------------------------------------------------------------
DRAFT_LOG_PREFIX = "Draft saved in browser"
_LOG_SKIP = {
    "name", "owner", "creation", "modified", "modified_by", "docstatus",
    "parent", "parentfield", "parenttype", "idx", "doctype",
}


def _fmt_val(v):
    return "(empty)" if v is None or v == "" else str(v)


def get_visit_logs(name):
    """Ek COM Visit Tracker record ki poori activity history (latest first). Details structured blocks me."""
    import json
    import re
    from html import unescape

    dt = "COM Visit Tracker"
    child_dt = "COM Visit and Comp Item"

    if not _is_admin(frappe.session.user):
        frappe.throw(frappe._("Not permitted"), frappe.PermissionError)
    if not frappe.db.exists(dt, name):
        frappe.throw(frappe._("Record not found"), frappe.DoesNotExistError)

    labels = {f.fieldname: (f.label or f.fieldname) for f in frappe.get_meta(dt).fields}
    c_labels = {f.fieldname: (f.label or f.fieldname) for f in frappe.get_meta(child_dt).fields}

    def who(user):
        full = frappe.utils.get_fullname(user)
        return {"user": user, "full_name": full if full and full != user else ""}

    def pairs(row):
        return [
            (c_labels.get(k, k), str(row[k]))
            for k in row
            if k not in _LOG_SKIP and row[k] not in (None, "") and str(row[k]) not in ("0", "0.0")
        ]

    def parse_draft_line(line):
        m = re.match(r"Row (\d+) - (.*)", line)
        if not m:
            return None
        kv = []
        for part in re.split(r", (?=[A-Z][\w ]*: )", m.group(2)):
            if ": " in part:
                k, v = part.split(": ", 1)
                kv.append((k, v))
        return int(m.group(1)), kv

    events = []

    # 1) Created
    created = frappe.db.get_value(dt, name, ["owner", "creation"], as_dict=True)
    events.append({"time": created.creation, "action": "Created", "details": [], **who(created.owner)})

    # 2) Version = har real save
    versions = frappe.get_all(
        "Version",
        filters={"ref_doctype": dt, "docname": name},
        fields=["creation", "owner", "data"],
        order_by="creation asc",
        limit_page_length=0,
    )
    for v in versions:
        try:
            data = json.loads(v.data or "{}")
        except Exception:
            data = {}
        blocks = []

        top = [
            (labels.get(ch[0], ch[0]), _fmt_val(ch[1]), _fmt_val(ch[2]))
            for ch in (data.get("changed") or []) if len(ch) >= 3
        ]
        if top:
            blocks.append({"t": "changes", "row": None, "changes": top})

        for item in data.get("added") or []:
            blocks.append({"t": "added", "fields": pairs(item[1])})

        for item in data.get("removed") or []:
            blocks.append({"t": "removed", "fields": pairs(item[1])})

        for item in data.get("row_changed") or []:
            idx, changes = item[1], item[3]
            ch_list = [
                (c_labels.get(c[0], c[0]), _fmt_val(c[1]), _fmt_val(c[2]))
                for c in (changes or []) if len(c) >= 3
            ]
            if ch_list:
                blocks.append({"t": "changes", "row": int(idx) + 1, "changes": ch_list})

        events.append({
            "time": v.creation,
            "action": "Updated",
            "details": blocks or [{"t": "text", "text": "Document saved"}],
            **who(v.owner),
        })

    # 3) Comments / attachments / draft saves
    comments = frappe.get_all(
        "Comment",
        filters={"reference_doctype": dt, "reference_name": name},
        fields=["creation", "owner", "comment_type", "content"],
        order_by="creation asc",
        limit_page_length=0,
    )
    for c in comments:
        content = c.content or ""
        lines = [
            unescape(frappe.utils.strip_html_tags(p)).strip()
            for p in re.split(r"<br\s*/?>", content)
        ]
        lines = [l for l in lines if l]
        if content.startswith(DRAFT_LOG_PREFIX):
            rows = [r for r in (parse_draft_line(l) for l in lines[1:]) if r]
            blocks = [{"t": "draft", "summary": lines[0] if lines else "", "rows": rows}]
            action = "Draft Saved"
        else:
            blocks = [{"t": "text", "text": l} for l in lines]
            action = c.comment_type or "Comment"
        events.append({"time": c.creation, "action": action, "details": blocks, **who(c.owner)})

    # 4) Views
    if frappe.db.exists("DocType", "View Log"):
        for vl in frappe.get_all(
            "View Log",
            filters={"reference_doctype": dt, "reference_name": name},
            fields=["creation", "viewed_by"],
            limit_page_length=0,
        ):
            events.append({
                "time": vl.creation,
                "action": "Viewed",
                "details": [{"t": "text", "text": "Record opened"}],
                **who(vl.viewed_by),
            })

    events.sort(key=lambda e: e["time"], reverse=True)
    return events


@frappe.whitelist()
def log_draft_save(name, rows=None):
    """Widget 'Save' (browser draft) ko bhi server pe log karo. Data/DB rows nahi badalta."""
    from html import escape

    dt = "COM Visit Tracker"
    if not frappe.has_permission(dt, "write", doc=name):
        frappe.throw(frappe._("Not permitted"), frappe.PermissionError)

    rows = frappe.parse_json(rows) or []
    labels = {f.fieldname: (f.label or f.fieldname) for f in frappe.get_meta("COM Visit and Comp Item").fields}

    lines = []
    for i, r in enumerate(rows, 1):
        parts = [f"{labels[k]}: {v}" for k, v in r.items() if k in labels and v not in (None, "")]
        lines.append(f"Row {i} - " + (", ".join(parts) or "(blank row)"))

    content = f"{DRAFT_LOG_PREFIX} (Final Submission Pending). Total rows: {len(rows)}"
    content += "".join("<br>" + escape(l) for l in lines)

    frappe.get_doc({
        "doctype": "Comment",
        "comment_type": "Info",
        "reference_doctype": dt,
        "reference_name": name,
        "content": content,
    }).insert(ignore_permissions=True)
