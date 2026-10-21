import frappe
from html import escape

from sahayog.branch_score_card.doctype.com_visit_tracker.com_visit_tracker import get_visit_logs

TEMPLATE = """<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Logs - __NAME__</title>
<style>
*{box-sizing:border-box}
body{margin:0;background:#fff;color:#1f2933;font:13px/1.5 -apple-system,"Segoe UI",Roboto,Arial,sans-serif}
.wrap{width:100%;padding:24px 32px}
.top{display:flex;justify-content:space-between;align-items:flex-end;gap:16px;flex-wrap:wrap;margin:0 0 14px}
.top h1{margin:0;font-size:18px;font-weight:600}
.sub{color:#8a94a0;font-size:12px;margin-top:2px}
.meta{display:flex;gap:18px;align-items:center;flex-wrap:wrap;color:#8a94a0;font-size:12px}
.pill{background:none;border:0;padding:0;color:#8a94a0;font-size:12px}
.pill b{color:#1f2933;font-weight:600;font-size:12px;margin-right:4px}
#q{border:0;border-bottom:1px solid #d5d9de;border-radius:0;padding:4px 2px;width:220px;font-size:13px;background:none;outline:none}
#q:focus{border-color:#1f2933}
.dl_bar{display:flex;gap:18px;align-items:center;flex-wrap:wrap;margin:0 0 6px;font-size:12px;color:#8a94a0}
.dl_bar input{border:0;border-bottom:1px solid #d5d9de;border-radius:0;padding:3px 2px;font-size:13px;background:none;color:#1f2933;outline:none}
.dl_bar input:focus{border-color:#1f2933}
.dl_bar button{border:1px solid #1f2933;background:none;color:#1f2933;border-radius:4px;padding:4px 12px;font-size:12px;cursor:pointer}
.dl_bar button:hover{background:#1f2933;color:#fff}
.card{background:none;border:0;border-radius:0;overflow:visible}
table{width:100%;border-collapse:collapse}
thead th{background:none;text-align:left;font-size:11px;font-weight:500;text-transform:uppercase;letter-spacing:.06em;color:#8a94a0;padding:10px 12px;border-bottom:1px solid #e5e7eb}
tbody td{padding:10px 12px;border-bottom:1px solid #f0f1f3;vertical-align:top;background:none}
tbody tr:hover{background:none}
tbody tr:last-child td{border-bottom:0}
.when{white-space:nowrap;width:140px;color:#6b7785}
.rec{width:190px;font-weight:500}
.who{width:150px}
.who b{display:block;font-weight:500}
.who span{color:#a0a8b1;font-size:12px}
.act{width:90px}
.badge{display:inline-block;background:none !important;border-radius:0;padding:0;font-size:12px;font-weight:500;color:#52606d}
.badge:before{content:"";display:inline-block;width:6px;height:6px;border-radius:50%;margin-right:6px;vertical-align:middle;background:#a0a8b1}
.badge.created{color:#067647}.badge.created:before{background:#067647}
.badge.updated{color:#1d56b8}.badge.updated:before{background:#1d56b8}
.badge.draft{color:#b45309}.badge.draft:before{background:#b45309}
.chgs{display:flex;flex-wrap:wrap;gap:2px 22px}
.ln{display:inline-block;background:none;border:0;border-radius:0;padding:0;margin:0 22px 2px 0;line-height:1.5}
.chgs .ln{margin:0}
.tb,.row{display:inline;background:none;border-radius:0;padding:0;min-width:0;margin-right:6px;color:#8a94a0;font-size:12px;font-weight:400}
.f{font-weight:500}
.old{color:#b42318;text-decoration:line-through}
.new{color:#067647;font-weight:500}
.arr{color:#a0a8b1;margin:0 4px}
.add{color:#067647;font-weight:500}
.rem{color:#b42318;font-weight:500}
.draft{color:#b45309;font-weight:500}
.muted{color:#a0a8b1}
.empty{padding:28px;text-align:center;color:#a0a8b1}
</style></head><body><div class="wrap">
<div class="top">
<div><h1>Activity Logs</h1><div class="sub">__NAME__</div></div>
<div class="meta"><div class="pill"><b>__N__</b>Events</div><div class="pill"><b>__U__</b>Users</div></div>
</div>
<div class="dl_bar" id="dl_bar" data-url="/api/method/sahayog.branch_score_card.doctype.com_visit_tracker.visit_logs_page.download_logs" data-name="__NAME__"><label>From <input type="date" id="dl_from"></label><label>To <input type="date" id="dl_to"></label><button type="button" id="dl_btn">Download Excel</button></div>
<div class="card"><table>
<thead><tr><th>Date / Time</th><th>User</th><th>Action</th><th>What changed</th></tr></thead>
<tbody>__ROWS__</tbody>
</table></div>
</div><script>
(function(){var b=document.getElementById('dl_bar');if(!b)return;
document.getElementById('dl_btn').onclick=function(){
var f=document.getElementById('dl_from').value,t=document.getElementById('dl_to').value;
if(f&&t&&f>t){alert('From Date cannot be after To Date');return;}
var p=[];if(b.dataset.name)p.push('name='+encodeURIComponent(b.dataset.name));
if(f)p.push('from_date='+f);if(t)p.push('to_date='+t);
window.location=b.dataset.url+(p.length?'?'+p.join('&'):'');};})();
</script>
</body></html>"""


def _fields(fields):
    return ", ".join("%s: %s" % (escape(l), escape(v)) for l, v in fields) or "blank row"


def _noise(old, new):
    empty = ("(empty)", "0", "0.0")
    return old in empty and new in empty


def _zero(v):
    try:
        return float(v) == 0
    except Exception:
        return False


def _lines(details):
    out = []
    for d in details:
        t = d.get("t")
        if t == "changes":
            row = "<span class='row'>Row %s</span>" % d["row"] if d.get("row") else "<span class='row'>Record</span>"
            for label, old, new in d["changes"]:
                if _noise(old, new):
                    continue
                out.append(
                    "<div class='ln'>%s<span class='f'>%s</span>: <span class='old'>%s</span>"
                    "<span class='arr'>&rarr;</span><span class='new'>%s</span></div>"
                    % (row, escape(label), escape(old), escape(new))
                )
        elif t == "added":
            out.append("<div class='ln'><span class='row'>Row</span><span class='add'>Added</span> &middot; %s</div>" % _fields(d["fields"]))
        elif t == "removed":
            out.append("<div class='ln'><span class='row'>Row</span><span class='rem'>Removed</span> &middot; %s</div>" % _fields(d["fields"]))
        elif t == "text":
            out.append("<div class='ln'>%s</div>" % escape(d.get("text", "")))
        elif t == "draft":
            for n, kv in d["rows"]:
                parts = ", ".join("%s: %s" % (escape(k), escape(v)) for k, v in kv if v and not _zero(v))
                out.append(
                    "<div class='ln'><span class='row'>Row %s</span><span class='draft'>Draft</span> &middot; %s</div>"
                    % (n, parts or "blank row")
                )
    return "".join(out)


@frappe.whitelist()
def view_logs_page(name):
    events = get_visit_logs(name)  # permission check yahin hota hai

    kind_map = {"Created": "created", "Updated": "updated", "Draft Saved": "draft"}
    rows = []
    shown = []
    for e in events:
        if e["action"] == "Viewed":
            continue
        if e["action"] == "Created":
            det = "<span class='muted'>Record created</span>"
        else:
            det = _lines(e["details"])
            if not det:
                continue
        shown.append(e)
        when = frappe.utils.format_datetime(e["time"], "dd-MM-yyyy HH:mm:ss")
        who = "<b>%s</b><span>%s</span>" % (escape(e.get("full_name") or e["user"] or ""), escape(e["user"] or ""))
        rows.append(
            "<tr><td class='when'>%s</td><td class='who'>%s</td>"
            "<td class='act'><span class='badge %s'>%s</span></td><td>%s</td></tr>"
            % (escape(when), who, kind_map.get(e["action"], "other"), escape(e["action"]), det)
        )

    html = (
        TEMPLATE.replace("__NAME__", escape(name))
        .replace("__N__", str(len(shown)))
        .replace("__U__", str(len({e["user"] for e in shown})))
        .replace("__ROWS__", "".join(rows))
    )

    frappe.local.response.filename = "logs.html"
    frappe.local.response.filecontent = html
    frappe.local.response.type = "download"
    frappe.local.response.display_content_as = "inline"


def _flat_lines(details):
    out = []
    for d in details:
        t = d.get("t")
        if t == "changes":
            row = "Row %s" % d["row"] if d.get("row") else "Record"
            for label, old, new in d["changes"]:
                if _noise(old, new):
                    continue
                out.append((row, label, old, new, ""))
        elif t in ("added", "removed"):
            txt = ", ".join("%s: %s" % (l, v) for l, v in d["fields"]) or "blank row"
            out.append(("Row", "Row Added" if t == "added" else "Row Removed", "", "", txt))
        elif t == "draft":
            for n, kv in d["rows"]:
                txt = ", ".join("%s: %s" % (k, v) for k, v in kv if v and not _zero(v)) or "blank row"
                out.append(("Row %s" % n, "Draft Saved", "", "", txt))
        elif t == "text":
            out.append(("", "", "", "", d.get("text", "")))
    return out


@frappe.whitelist()
def download_logs(name, from_date=None, to_date=None):
    import io
    from frappe.utils import getdate
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment

    events = get_visit_logs(name)  # permission check yahin hota hai
    fd = getdate(from_date) if from_date else None
    td = getdate(to_date) if to_date else None
    if fd and td and fd > td:
        frappe.throw(frappe._("From Date cannot be after To Date"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Logs"
    ws.append(["Date / Time", "Record", "User", "Name", "Action", "Row", "Field", "Old Value", "New Value", "Details"])

    for e in sorted(events, key=lambda x: x["time"], reverse=True):
        if e["action"] == "Viewed":
            continue
        d = getdate(e["time"])
        if (fd and d < fd) or (td and d > td):
            continue
        when = frappe.utils.format_datetime(e["time"], "dd-MM-yyyy HH:mm:ss")
        if e["action"] == "Created":
            lines = [("Record", "", "", "", "Record created")]
        else:
            lines = _flat_lines(e["details"])
        for ln in lines:
            ws.append([when, name, e["user"], e.get("full_name") or "", e["action"], *ln])

    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for col in ws.columns:
        w = max(len(str(c.value)) if c.value is not None else 0 for c in col[:300])
        ws.column_dimensions[col[0].column_letter].width = min(max(w + 2, 12), 60)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)

    buf = io.BytesIO()
    wb.save(buf)
    frappe.response["filename"] = "COM_Visit_Tracker_Logs_%s.xlsx" % frappe.utils.today()
    frappe.response["filecontent"] = buf.getvalue()
    frappe.response["type"] = "binary"
