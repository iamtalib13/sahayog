import json
from html import escape

import frappe
from sahayog.branch_score_card.audit_log_utils import get_deleted_record_audit_logs, audit_details, audit_details_html, audit_details_text

DT = "Miscellaneous"
TABLES = {
    "account_opening_error": "Account Opening Error",
    "bank_reconciliation_discrepancy": "Bank Reconciliation Discrepancy",
}
SKIP = {
    "name", "owner", "creation", "modified", "modified_by", "docstatus",
    "parent", "parentfield", "parenttype", "idx", "doctype", "__unsaved",
}
EMPTY = ("(empty)", "0", "0.0")

TEMPLATE = """<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Miscellaneous - Logs</title>
<style>
*{box-sizing:border-box}
body{margin:0;background:#fff;color:#1f2933;font:13px/1.5 -apple-system,"Segoe UI",Roboto,Arial,sans-serif}
.wrap{width:100%;padding:24px 32px}
.top{display:flex;justify-content:space-between;align-items:flex-end;gap:16px;flex-wrap:wrap;margin:0 0 14px}
.top h1{margin:0;font-size:18px;font-weight:600}
.sub{color:#8a94a0;font-size:12px;margin-top:2px}
.meta{display:flex;gap:18px;align-items:center;flex-wrap:wrap;color:#8a94a0;font-size:12px}
.pill{background:none;border:0;padding:0;color:#8a94a0;font-size:12px}
.pill b{color:#1f2933;font-weight:600;font-size:12px;margin-right:4px}
#q{border:0;border-bottom:1px solid #d5d9de;border-radius:0;padding:4px 2px;width:220px;font-size:13px;background:none;outline:none}
#q:focus{border-color:#1f2933}
.dl_bar{display:flex;gap:18px;align-items:center;flex-wrap:wrap;margin:0 0 6px;font-size:12px;color:#8a94a0}
.dl_bar input{border:0;border-bottom:1px solid #d5d9de;border-radius:0;padding:3px 2px;font-size:13px;background:none;color:#1f2933;outline:none}
.dl_bar input:focus{border-color:#1f2933}
.dl_bar button{border:1px solid #1f2933;background:none;color:#1f2933;border-radius:4px;padding:4px 12px;font-size:12px;cursor:pointer}
.dl_bar button:hover{background:#1f2933;color:#fff}
.card{background:none;border:0;border-radius:0;overflow:visible}
table{width:100%;border-collapse:collapse}
thead th{background:none;text-align:left;font-size:11px;font-weight:500;text-transform:uppercase;letter-spacing:.06em;color:#8a94a0;padding:10px 12px;border-bottom:1px solid #e5e7eb}
tbody td{padding:10px 12px;border-bottom:1px solid #f0f1f3;vertical-align:top;background:none}
tbody tr:hover{background:none}
tbody tr:last-child td{border-bottom:0}
.when{white-space:nowrap;width:140px;color:#6b7785}
.rec{width:190px;font-weight:500}
.who{width:150px}
.who b{display:block;font-weight:500}
.who span{color:#a0a8b1;font-size:12px}
.act{width:90px}
.badge{display:inline-block;background:none !important;border-radius:0;padding:0;font-size:12px;font-weight:500;color:#52606d}
.badge:before{content:"";display:inline-block;width:6px;height:6px;border-radius:50%;margin-right:6px;vertical-align:middle;background:#a0a8b1}
.badge.created{color:#067647}.badge.created:before{background:#067647}
.badge.updated{color:#1d56b8}.badge.updated:before{background:#1d56b8}
.badge.draft{color:#b45309}.badge.draft:before{background:#b45309}
.chgs{display:flex;flex-wrap:wrap;gap:2px 22px}
.ln{display:inline-block;background:none;border:0;border-radius:0;padding:0;margin:0 22px 2px 0;line-height:1.5}
.chgs .ln{margin:0}
.tb,.row{display:inline;background:none;border-radius:0;padding:0;min-width:0;margin-right:6px;color:#8a94a0;font-size:12px;font-weight:400}
.f{font-weight:500}
.old{color:#b42318;text-decoration:line-through}
.new{color:#067647;font-weight:500}
.arr{color:#a0a8b1;margin:0 4px}
.add{color:#067647;font-weight:500}
.rem{color:#b42318;font-weight:500}
.draft{color:#b45309;font-weight:500}
.muted{color:#a0a8b1}
.empty{padding:28px;text-align:center;color:#a0a8b1}
</style></head><body><div class="wrap">
<div class="top">
<div><h1>Miscellaneous - Activity Logs</h1><div class="sub">All records &middot; latest first</div></div>
<div class="meta">
<input id="q" type="text" placeholder="Search branch / SOL / user...">
<div class="pill"><b>__N__</b>Events</div>
<div class="pill"><b>__B__</b>Records</div>
<div class="pill"><b>__U__</b>Users</div>
</div></div>
<div class="dl_bar" id="dl_bar" data-url="/api/method/sahayog.branch_score_card.doctype.miscellaneous.miscellaneous_logs_page.download_logs"><label>From <input type="date" id="dl_from"></label><label>To <input type="date" id="dl_to"></label><button type="button" id="dl_btn">Download Excel</button></div>
<div class="card"><table>
<thead><tr><th>Date / Time</th><th>Record</th><th>User</th><th>Action</th><th>What changed</th></tr></thead>
<tbody id="tb">__ROWS__</tbody>
</table><div class="empty" id="empty">No matching logs</div></div>
</div>
<script>
var q=document.getElementById('q'),rows=document.querySelectorAll('#tb tr'),em=document.getElementById('empty');
q.addEventListener('input',function(){var v=q.value.toLowerCase().trim(),n=0;
rows.forEach(function(r){var ok=!v||r.getAttribute('data-s').indexOf(v)>-1;r.style.display=ok?'':'none';if(ok)n++;});
em.style.display=n?'none':'block';});
</script><script>
(function(){var b=document.getElementById('dl_bar');if(!b)return;
document.getElementById('dl_btn').onclick=function(){
var f=document.getElementById('dl_from').value,t=document.getElementById('dl_to').value;
if(f&&t&&f>t){alert('From Date cannot be after To Date');return;}
var p=[];if(b.dataset.name)p.push('name='+encodeURIComponent(b.dataset.name));
if(f)p.push('from_date='+f);if(t)p.push('to_date='+t);
window.location=b.dataset.url+(p.length?'?'+p.join('&'):'');};})();
</script>
</body></html>"""


def _can_view():
    u = frappe.session.user
    return u == "Administrator" or bool(
        {"System Manager", "Branch Scorecard Editor"} & set(frappe.get_roles(u))
    )


def _v(x):
    return "(empty)" if x is None or x == "" else str(x)


@frappe.whitelist()
def view_logs_page():
    if not _can_view():
        frappe.throw(frappe._("Not permitted"), frappe.PermissionError)

    meta_cache = {}

    def labels(doctype):
        if doctype not in meta_cache:
            meta_cache[doctype] = {
                f.fieldname: (f.label or f.fieldname) for f in frappe.get_meta(doctype).fields
            }
        return meta_cache[doctype]

    parent_labels = labels(DT)
    child_dt = {f.fieldname: f.options for f in frappe.get_meta(DT).fields if f.fieldname in TABLES}
    names = {}

    def fullname(u):
        if u not in names:
            fn = frappe.utils.get_fullname(u)
            names[u] = fn if fn and fn != u else ""
        return names[u]

    def clabel(table, field):
        return labels(child_dt[table]).get(field, field) if table in child_dt else field

    def pairs(table, row):
        return ", ".join(
            "%s: %s" % (escape(clabel(table, k)), escape(str(v)))
            for k, v in row.items()
            if k not in SKIP and v not in (None, "") and str(v) not in ("0", "0.0")
        ) or "blank row"

    events = []

    for d in frappe.get_all(DT, fields=["name", "owner", "creation"], limit_page_length=0):
        events.append({
            "time": d.creation, "record": d.name, "user": d.owner, "action": "Created",
            "html": "<span class='muted'>Record created</span>",
        })

    versions = frappe.get_all(
        "Version",
        filters={"ref_doctype": DT},
        fields=["creation", "owner", "docname", "data"],
        order_by="creation desc",
        limit_page_length=3000,
    )
    for v in versions:
        try:
            data = json.loads(v.data or "{}")
        except Exception:
            continue
        out = []

        deleted_snapshot = data.get("deleted")
        if isinstance(deleted_snapshot, dict):
            snapshot_html = escape(
                json.dumps(deleted_snapshot, ensure_ascii=False, default=str)
            )
            out.append(
                "<div class='ln'><span class='tb'>Deleted snapshot</span>"
                "<span class='f'>%s</span></div>" % snapshot_html
            )

        for ch in data.get("changed") or []:
            if len(ch) >= 3 and not (_v(ch[1]) in EMPTY and _v(ch[2]) in EMPTY):
                out.append(
                    "<div class='ln'><span class='tb'>Record</span><span class='f'>%s</span>: "
                    "<span class='old'>%s</span><span class='arr'>&rarr;</span><span class='new'>%s</span></div>"
                    % (escape(parent_labels.get(ch[0], ch[0])), escape(_v(ch[1])), escape(_v(ch[2])))
                )

        for kind, cls, word in (("added", "add", "Added"), ("removed", "rem", "Removed")):
            for item in data.get(kind) or []:
                if len(item) >= 2 and isinstance(item[1], dict):
                    t = item[0]
                    out.append(
                        "<div class='ln'><span class='tb'>%s</span><span class='%s'>Row %s</span> &middot; %s</div>"
                        % (escape(TABLES.get(t, t)), cls, word, pairs(t, item[1]))
                    )

        for item in data.get("row_changed") or []:
            if len(item) < 4:
                continue
            t, idx, _rn, changes = item[0], item[1], item[2], item[3]
            for c in changes or []:
                if len(c) >= 3 and not (_v(c[1]) in EMPTY and _v(c[2]) in EMPTY):
                    out.append(
                        "<div class='ln'><span class='tb'>%s &middot; Row %s</span><span class='f'>%s</span>: "
                        "<span class='old'>%s</span><span class='arr'>&rarr;</span><span class='new'>%s</span></div>"
                        % (escape(TABLES.get(t, t)), int(idx) + 1, escape(clabel(t, c[0])),
                           escape(_v(c[1])), escape(_v(c[2])))
                    )

        if out:
            events.append({
                "time": v.creation, "record": v.docname, "user": v.owner,
                "action": "Deleted" if deleted_snapshot else "Updated",
                "html": "<div class='chgs'>" + "".join(out) + "</div>",
            })

    permanent_rows = get_deleted_record_audit_logs(DT)
    deleted_names = {row.record_name for row in permanent_rows}

    # For deleted records, use permanent history instead of duplicate Version/live events.
    events = [event for event in events if event["record"] not in deleted_names]

    for audit in permanent_rows:
        details = audit_details_html(audit)
        events.append({
            "time": audit.event_time,
            "record": audit.record_name,
            "user": audit.user,
            "action": audit.action,
            "html": (
                "<div class='chgs'><div class='ln'>"
                "<span class='tb'>Permanent audit · Source: %s</span>"
                "<span class='f'>%s</span></div></div>"
            ) % (escape(audit.source or "Manual"), details),
        })

    events.sort(key=lambda e: e["time"], reverse=True)

    rows = []
    for e in events:
        fn = fullname(e["user"])
        when = frappe.utils.format_datetime(e["time"], "dd-MM-yyyy HH:mm:ss")
        search = escape(("%s %s %s" % (e["record"], e["user"], fn)).lower(), quote=True)
        rows.append(
            "<tr data-s=\"%s\"><td class='when'>%s</td><td class='rec'>%s</td>"
            "<td class='who'><b>%s</b><span>%s</span></td>"
            "<td class='act'><span class='badge %s'>%s</span></td><td>%s</td></tr>"
            % (search, escape(when), escape(e["record"]), escape(fn or e["user"] or ""),
               escape(e["user"] or ""), e["action"].lower(), e["action"], e["html"])
        )

    html = (
        TEMPLATE.replace("__N__", str(len(events)))
        .replace("__B__", str(len({e["record"] for e in events})))
        .replace("__U__", str(len({e["user"] for e in events})))
        .replace("__ROWS__", "".join(rows))
    )

    frappe.local.response.filename = "miscellaneous_logs.html"
    frappe.local.response.filecontent = html
    frappe.local.response.type = "download"
    frappe.local.response.display_content_as = "inline"


def _lm(doctype):
    return {f.fieldname: (f.label or f.fieldname) for f in frappe.get_meta(doctype).fields}


@frappe.whitelist()
def download_logs(from_date=None, to_date=None):
    import io
    from frappe.utils import getdate, add_days, format_datetime
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment

    if not _can_view():
        frappe.throw(frappe._("Not permitted"), frappe.PermissionError)

    fd = getdate(from_date) if from_date else None
    td = getdate(to_date) if to_date else None
    if fd and td and fd > td:
        frappe.throw(frappe._("From Date cannot be after To Date"))

    date_f = []
    if fd:
        date_f.append(["creation", ">=", fd])
    if td:
        date_f.append(["creation", "<", add_days(td, 1)])

    parent_l = _lm(DT)
    child_dt = {f.fieldname: f.options for f in frappe.get_meta(DT).fields if f.fieldname in TABLES}
    child_l = {t: _lm(o) for t, o in child_dt.items()}
    names = {}

    def full(u):
        if u not in names:
            fn = frappe.utils.get_fullname(u)
            names[u] = fn if fn and fn != u else ""
        return names[u]

    def pairs(t, row):
        return ", ".join(
            "%s: %s" % (child_l.get(t, {}).get(k, k), v)
            for k, v in row.items()
            if k not in SKIP and v not in (None, "") and str(v) not in ("0", "0.0")
        ) or "blank row"

    # (time, record, user, action, table, row, field, old, new, details)
    items = []

    for d in frappe.get_all(DT, filters=date_f, fields=["name", "owner", "creation"], limit_page_length=0):
        items.append((d.creation, d.name, d.owner, "Created", "", "", "", "", "", "Record created"))

    versions = frappe.get_all(
        "Version",
        filters=[["ref_doctype", "=", DT]] + date_f,
        fields=["creation", "owner", "docname", "data"],
        limit_page_length=0,
    )
    for v in versions:
        try:
            data = json.loads(v.data or "{}")
        except Exception:
            continue
        deleted_snapshot = data.get("deleted")
        if isinstance(deleted_snapshot, dict):
            items.append((
                v.creation, v.docname, v.owner, "Deleted",
                "Record", "", "Deleted snapshot", "", "",
                json.dumps(deleted_snapshot, ensure_ascii=False, default=str),
            ))
            continue

        base = (v.creation, v.docname, v.owner, "Updated")

        for ch in data.get("changed") or []:
            if len(ch) >= 3 and not (_v(ch[1]) in EMPTY and _v(ch[2]) in EMPTY):
                items.append(base + ("Record", "", parent_l.get(ch[0], ch[0]), _v(ch[1]), _v(ch[2]), ""))

        for kind, word in (("added", "Row Added"), ("removed", "Row Removed")):
            for item in data.get(kind) or []:
                if len(item) >= 2 and isinstance(item[1], dict):
                    t = item[0]
                    items.append(base + (TABLES.get(t, t), "", word, "", "", pairs(t, item[1])))

        for item in data.get("row_changed") or []:
            if len(item) < 4:
                continue
            t, idx, changes = item[0], item[1], item[3]
            for c in changes or []:
                if len(c) >= 3 and not (_v(c[1]) in EMPTY and _v(c[2]) in EMPTY):
                    items.append(base + (
                        TABLES.get(t, t), int(idx) + 1,
                        child_l.get(t, {}).get(c[0], c[0]), _v(c[1]), _v(c[2]), "",
                    ))

    permanent_rows = get_deleted_record_audit_logs(DT)
    deleted_names = {row.record_name for row in permanent_rows}

    # Replace legacy events for deleted records with their permanent audit history.
    items = [item for item in items if item[1] not in deleted_names]

    for audit in permanent_rows:
        if fd and frappe.utils.getdate(audit.event_time) < fd:
            continue
        if td and frappe.utils.getdate(audit.event_time) > td:
            continue
        items.append((
            audit.event_time,
            audit.record_name,
            audit.user,
            audit.action,
            "Record",
            "",
            "Permanent audit details",
            "",
            "",
            audit_details_text(audit),
        ))

    items.sort(key=lambda x: x[0], reverse=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Logs"
    ws.append(["Date / Time", "Record", "User", "Name", "Action", "Table", "Row", "Field", "Old Value", "New Value", "Details"])
    for it in items:
        ws.append([
            format_datetime(it[0], "dd-MM-yyyy HH:mm:ss"), it[1], it[2], full(it[2]), it[3],
            it[4], it[5], it[6], it[7], it[8], it[9],
        ])

    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        w = max(len(str(c.value)) if c.value is not None else 0 for c in col[:300])
        ws.column_dimensions[col[0].column_letter].width = min(max(w + 2, 12), 60)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)

    buf = io.BytesIO()
    wb.save(buf)
    frappe.response["filename"] = "Miscellaneous_Logs_%s.xlsx" % frappe.utils.today()
    frappe.response["filecontent"] = buf.getvalue()
    frappe.response["type"] = "binary"
